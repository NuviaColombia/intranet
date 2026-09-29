"""Lógica de negocio de Caja menor Nuvia: recibos, legalización (FM), arqueos e importación desde la
hoja de Google de la app anterior (Apps Script)."""
import json
from datetime import datetime, date, timedelta
from sqlalchemy import func
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload
from .models import Empleado
from .models_caja import (CajaMenor, CajaAcceso, CajaAutorizador, CajaSupervisor, CajaRecibo, CajaFM, CajaArqueo,
                          CajaObservacion)
from .formato import nombre_propio

MONEDAS = [50, 100, 200, 500, 1000]
BILLETES = [2000, 5000, 10000, 20000, 50000, 100000]
DENOMINACIONES = MONEDAS + BILLETES

# Cajas con las que arranca el módulo (mismos prefijos y consecutivos que la app anterior).
CAJAS_INICIALES = [
    {"nombre": "CONTABILIDAD", "prefijo": "CONT", "consecutivo_inicial": 92, "fm_siguiente": 28, "icono": "📒"},
    {"nombre": "MANTENIMIENTO", "prefijo": "MANT", "consecutivo_inicial": 25, "fm_siguiente": 2, "icono": "🔧"},
]


# ---------------- Utilidades ----------------

def hora_colombia(dt: datetime | None) -> str:
    """Las fechas se guardan en UTC; Colombia es UTC-5 (sin horario de verano)."""
    return (dt - timedelta(hours=5)).strftime("%Y-%m-%d %H:%M") if dt else ""


def _unidades(n: int) -> str:
    return ["", "un", "dos", "tres", "cuatro", "cinco", "seis", "siete", "ocho", "nueve"][n]


def _decenas(n: int) -> str:
    d, u = divmod(n, 10)
    if d == 0:
        return _unidades(u)
    if d == 1:
        especiales = ["diez", "once", "doce", "trece", "catorce", "quince"]
        return especiales[u] if u <= 5 else "dieci" + _unidades(u)
    if d == 2:
        return "veinte" if u == 0 else "veinti" + _unidades(u)
    nombres = {3: "treinta", 4: "cuarenta", 5: "cincuenta", 6: "sesenta", 7: "setenta", 8: "ochenta", 9: "noventa"}
    return nombres[d] + (" y " + _unidades(u) if u else "")


def _centenas(n: int) -> str:
    c, resto = divmod(n, 100)
    if c == 0:
        return _decenas(resto)
    if c == 1:
        return "cien" if resto == 0 else "ciento " + _decenas(resto)
    nombres = {2: "doscientos", 3: "trescientos", 4: "cuatrocientos", 5: "quinientos", 6: "seiscientos",
               7: "setecientos", 8: "ochocientos", 9: "novecientos"}
    return nombres[c] + (" " + _decenas(resto) if resto else "")


def numero_a_letras(valor: float) -> str:
    """Valor en letras para el recibo, ej. 50000 -> "CINCUENTA MIL COP" (corrige el error de la app
    anterior, que escribía "sesenta" para los cincuenta)."""
    n = int(valor or 0)
    if n <= 0:
        return ""
    millones, resto = divmod(n, 1_000_000)
    miles, unidades = divmod(resto, 1000)
    partes = []
    if millones:
        partes.append("un millón" if millones == 1 else f"{_centenas(millones)} millones")
    if miles:
        partes.append("un mil" if miles == 1 else f"{_centenas(miles)} mil")  # como la app anterior
    if unidades:
        partes.append(_centenas(unidades))
    return f"{' '.join(partes)} COP".upper()


def _desglose_valido(desglose: dict) -> dict:
    limpio = {}
    for d in DENOMINACIONES:
        try:
            cant = int(desglose.get(str(d), desglose.get(d, 0)) or 0)
        except (TypeError, ValueError):
            cant = 0
        limpio[str(d)] = max(cant, 0)
    return limpio


def total_desglose(desglose: dict) -> float:
    return float(sum(int(k) * v for k, v in _desglose_valido(desglose).items()))


# ---------------- Cajas y permisos ----------------

def asegurar_cajas_iniciales(db: Session) -> None:
    if db.query(CajaMenor).count() == 0:
        for i, c in enumerate(CAJAS_INICIALES, start=1):
            db.add(CajaMenor(orden=i, **c))
        db.commit()


def es_admin(user: Empleado) -> bool:
    return user.rol in ("admin", "superadmin")


def cajas_de_usuario(db: Session, user: Empleado, incluir_inactivas: bool = False) -> list[CajaMenor]:
    q = db.query(CajaMenor)
    if not incluir_inactivas:
        q = q.filter(CajaMenor.activo == 1)
    if not es_admin(user):
        ids = [a.caja_id for a in db.query(CajaAcceso).filter(CajaAcceso.empleado_id == user.id)]
        q = q.filter(CajaMenor.id.in_(ids or [0]))
    return q.order_by(CajaMenor.orden, CajaMenor.nombre).all()


def ids_cajas_que_autoriza(db: Session, user: Empleado) -> set[int]:
    return {a.caja_id for a in db.query(CajaAutorizador).filter(CajaAutorizador.empleado_id == user.id)}


def ids_cajas_que_supervisa(db: Session, user: Empleado) -> set[int]:
    return {a.caja_id for a in db.query(CajaSupervisor).filter(CajaSupervisor.empleado_id == user.id)}


def autoriza_en_caja(db: Session, user: Empleado, caja: CajaMenor) -> bool:
    """Quien autoriza recibos o supervisa FMs en la caja (o tiene documentos de ella asignados) puede entrar
    a firmar, aunque no tenga acceso a la caja."""
    if caja.id in ids_cajas_que_autoriza(db, user) or caja.id in ids_cajas_que_supervisa(db, user):
        return True
    if db.query(CajaFM.id).filter(CajaFM.caja_id == caja.id, CajaFM.elaborado_por_id == user.id).first():
        return True
    if db.query(CajaRecibo.id).filter(CajaRecibo.caja_id == caja.id, CajaRecibo.autorizado_por_id == user.id).first():
        return True
    return db.query(CajaFM.id).filter(CajaFM.caja_id == caja.id, CajaFM.supervisado_por_id == user.id).first() is not None


def pendientes_de_firma(db: Session, user: Empleado) -> dict[int, int]:
    """{caja_id: recibos + FMs que `user` debe firmar} (para el número rojo de cada caja)."""
    total: dict[int, int] = {}
    filas = (db.query(CajaRecibo.caja_id, func.count(CajaRecibo.id))
             .filter(CajaRecibo.autorizado_por_id == user.id, CajaRecibo.firmado_en.is_(None), CajaRecibo.estado != "ANULADO")
             .group_by(CajaRecibo.caja_id).all())
    filas += (db.query(CajaFM.caja_id, func.count(CajaFM.id))
              .filter(CajaFM.elaborado_por_id == user.id, CajaFM.elaborado_en.is_(None), CajaFM.estado == "VIGENTE")
              .group_by(CajaFM.caja_id).all())
    filas += (db.query(CajaFM.caja_id, func.count(CajaFM.id))
              .filter(CajaFM.supervisado_por_id == user.id, CajaFM.supervisado_en.is_(None), CajaFM.estado == "VIGENTE",
                      (CajaFM.elaborado_por_id.is_(None)) | (CajaFM.elaborado_en.isnot(None)))
              .group_by(CajaFM.caja_id).all())
    for caja_id, n in filas:
        total[caja_id] = total.get(caja_id, 0) + n
    return total


def responsable_de_caja(db: Session, caja: CajaMenor) -> Empleado | None:
    """Empleado activo que es el responsable de la caja (Parámetros › Cajas); firma el "Elaborado por" del FM."""
    if caja.responsable_id:
        e = db.get(Empleado, caja.responsable_id)
        return e if e and e.activo else None
    nombre = (caja.responsable or "").strip().upper()
    if not nombre:
        return None
    for e in db.query(Empleado).filter(Empleado.activo == 1).all():
        if e.nombre_completo.strip().upper() == nombre:
            return e
    return None


def es_firmante_de_caja(db: Session, user: Empleado, caja: CajaMenor) -> bool:
    """Quien está en Autorizan o Supervisan de la caja: solo consulta y firma (no crea, edita ni anula).
    Los administradores conservan su acceso completo."""
    if es_admin(user):
        return False
    return caja.id in ids_cajas_que_autoriza(db, user) or caja.id in ids_cajas_que_supervisa(db, user)


def _serializar_observaciones(obs) -> tuple[list[dict], dict | None]:
    lista = [{"id": o.id, "texto": o.texto, "autor": nombre_propio(o.autor.nombre_completo) if o.autor else "",
              "creadoEn": hora_colombia(o.creado_en), "atendidaEn": hora_colombia(o.atendida_en),
              "atendidaPor": nombre_propio(o.atendida_por.nombre_completo) if o.atendida_por else ""} for o in obs]
    pendiente = next((o for o in reversed(lista) if not o["atendidaEn"]), None)
    return lista, pendiente


def _atender_observaciones(doc, user: Empleado) -> str:
    """Marca como atendidas las observaciones pendientes del documento; devuelve el texto de la última."""
    ultima = ""
    for o in getattr(doc, "observaciones", []) or []:
        if not o.atendida_en:
            o.atendida_en, o.atendida_por_id, ultima = datetime.utcnow(), user.id, o.texto
    return ultima


def hoy_colombia() -> date:
    return (datetime.utcnow() - timedelta(hours=5)).date()


def puede_usar_caja(db: Session, user: Empleado, caja: CajaMenor) -> bool:
    if es_admin(user):
        return True
    return db.query(CajaAcceso).filter(CajaAcceso.caja_id == caja.id, CajaAcceso.empleado_id == user.id).first() is not None


# ---------------- Recibos ----------------

def siguiente_consecutivo(db: Session, caja: CajaMenor) -> int:
    maximo = db.query(func.max(CajaRecibo.consecutivo)).filter(CajaRecibo.caja_id == caja.id).scalar()
    return max((maximo or 0) + 1, caja.consecutivo_inicial or 1)


def serializar_recibo(r: CajaRecibo) -> dict:
    return {
        "id": r.id, "consecutivo": r.consecutivo, "noRecibo": f"{r.caja.prefijo}-{r.consecutivo:04d}",
        "ciudad": r.ciudad, "fecha": r.fecha.isoformat() if r.fecha else "", "identificacion": r.identificacion,
        "pagadoA": r.pagado_a, "valor": r.valor, "valorLetras": r.valor_letras, "concepto": r.concepto,
        "factura": r.numero_factura, "anexo": bool(r.anexo), "autorizadoPor": r.autorizado_por or "", "estado": r.estado,
        "autorizadoPorId": r.autorizado_por_id, "firmaEmail": r.firma_email or "", "firmadoEn": hora_colombia(r.firmado_en),
        "avisoOk": r.aviso_ok, "avisoEn": hora_colombia(r.aviso_en),
        "observaciones": _serializar_observaciones(r.observaciones)[0],
        "observacionPendiente": _serializar_observaciones(r.observaciones)[1],
        "caja": nombre_propio(r.caja.nombre),
        "fm": f"FM{r.fm.numero}" if r.fm else "", "fmId": r.fm_id,
        "creadoPor": nombre_propio(r.creado_por.nombre_completo) if r.creado_por else r.creado_por_texto,
        "creadoEn": hora_colombia(r.creado_en),
        "anuladoPor": nombre_propio(r.anulado_por.nombre_completo) if r.anulado_por else "",
        "anuladoEn": hora_colombia(r.anulado_en), "motivoAnulacion": r.motivo_anulacion or "",
    }


def supervisores_de_caja(db: Session, caja: CajaMenor) -> list[str]:
    """Nombres (en mayúsculas) de quienes supervisan (visto bueno) los FM de la caja."""
    filas = (db.query(Empleado).join(CajaSupervisor, CajaSupervisor.empleado_id == Empleado.id)
             .filter(CajaSupervisor.caja_id == caja.id, Empleado.activo == 1)
             .order_by(Empleado.nombres, Empleado.apellidos).all())
    return [e.nombre_completo.strip().upper() for e in filas]


def autorizadores_de_caja(db: Session, caja: CajaMenor) -> list[str]:
    """Nombres (en mayúsculas, como se guardan en el recibo) de quienes autorizan los recibos de la caja."""
    filas = (db.query(Empleado).join(CajaAutorizador, CajaAutorizador.empleado_id == Empleado.id)
             .filter(CajaAutorizador.caja_id == caja.id, Empleado.activo == 1)
             .order_by(Empleado.nombres, Empleado.apellidos).all())
    return [e.nombre_completo.strip().upper() for e in filas]


def _resolver_autoriza(db: Session, caja: CajaMenor, datos: dict, actual: CajaRecibo | None = None,
                       creador_id: int | None = None):
    """Si la caja tiene autorizadores, el recibo debe llevar uno de ellos (o conservar el que ya tenía).
    Devuelve (empleado que autoriza o None, mensaje de error o None)."""
    elegido = str(datos.get("autorizado_por") or "").strip().upper()
    filas = (db.query(Empleado).join(CajaAutorizador, CajaAutorizador.empleado_id == Empleado.id)
             .filter(CajaAutorizador.caja_id == caja.id, Empleado.activo == 1).all())
    for e in filas:
        if e.nombre_completo.strip().upper() == elegido:
            if creador_id and e.id == creador_id:
                return None, "Quien registra el recibo no puede autorizarlo. Elige a otra persona en «Autorizado por»."
            return e, None
    if actual and elegido and elegido == (actual.autorizado_por or ""):
        return (db.get(Empleado, actual.autorizado_por_id) if actual.autorizado_por_id else None), None
    if filas and creador_id and all(e.id == creador_id for e in filas):
        return None, ("Eres el único autorizador de esta caja y no puedes firmar tu propio recibo. "
                      "Pide que agreguen a otra persona en Parámetros › Autorizan.")
    if filas:
        return None, "Elige quién autoriza el recibo."
    return None, None


def _confirmaciones(db: Session, caja: CajaMenor, limpios: dict, datos: dict, actual: CajaRecibo | None = None) -> dict | None:
    """Casos que no son error pero hay que confirmar: fecha de hace más de 30 días y factura ya registrada
    del mismo tercero. Devuelve {"codigo", "mensaje"} si falta confirmar."""
    confirmados = set(datos.get("confirmar") or [])
    if "fecha_antigua" not in confirmados and limpios["fecha"] < hoy_colombia() - timedelta(days=30):
        return {"codigo": "fecha_antigua",
                "mensaje": f"La fecha del recibo ({limpios['fecha'].strftime('%d/%m/%Y')}) es de hace más de 30 días. ¿Es correcta?"}
    factura = limpios.get("numero_factura") or ""
    if factura and "factura_repetida" not in confirmados:
        q = (db.query(CajaRecibo).filter(CajaRecibo.numero_factura == factura, CajaRecibo.identificacion == limpios["identificacion"],
                                         CajaRecibo.estado != "ANULADO"))
        if actual:
            q = q.filter(CajaRecibo.id != actual.id)
        otro = q.first()
        if otro:
            return {"codigo": "factura_repetida",
                    "mensaje": f"La factura {factura} de este tercero ya está registrada en el recibo "
                               f"{otro.caja.prefijo}-{otro.consecutivo:04d} (caja {nombre_propio(otro.caja.nombre)}). ¿Guardar de todas formas?"}
    return None


def _datos_recibo(caja: CajaMenor, datos: dict) -> dict | str:
    try:
        fecha = date.fromisoformat(str(datos.get("fecha") or ""))
    except ValueError:
        return "Fecha inválida."
    try:
        valor = round(float(datos.get("valor") or 0), 2)
    except (TypeError, ValueError):
        return "Valor inválido."
    if valor <= 0:
        return "El valor debe ser mayor que cero."
    if fecha > hoy_colombia():
        return "La fecha del recibo no puede ser futura."
    campos = {"identificacion": "No. identificación", "pagado_a": "Pagado a", "concepto": "Concepto"}
    limpios = {k: str(datos.get(k) or "").strip().upper() for k in ("identificacion", "pagado_a", "concepto")}
    for k, nombre in campos.items():
        if not limpios[k]:
            return f"El campo «{nombre}» es obligatorio."
    return {"ciudad": str(datos.get("ciudad") or caja.ciudad).strip().upper(), "fecha": fecha, "valor": valor,
            "valor_letras": numero_a_letras(valor), "numero_factura": str(datos.get("numero_factura") or "").strip().upper(),
            "anexo": 1 if datos.get("anexo") else 0,
            "autorizado_por": str(datos.get("autorizado_por") or "").strip().upper(), **limpios}


def crear_recibo(db: Session, caja: CajaMenor, user: Empleado, datos: dict) -> CajaRecibo | str:
    limpios = _datos_recibo(caja, datos)
    if isinstance(limpios, str):
        return limpios
    autoriza, error = _resolver_autoriza(db, caja, datos, creador_id=user.id)
    if error:
        return error
    aviso = _confirmaciones(db, caja, limpios, datos)
    if aviso:
        return aviso
    for _ in range(3):  # si dos personas guardan a la vez, el segundo toma el siguiente consecutivo
        consecutivo = siguiente_consecutivo(db, caja)
        r = CajaRecibo(caja_id=caja.id, consecutivo=consecutivo, creado_por_id=user.id,
                       autorizado_por_id=autoriza.id if autoriza else None, **limpios)
        if not r.numero_factura:
            r.numero_factura = f"{caja.prefijo}-{consecutivo:04d}"  # igual que la app anterior
        db.add(r)
        try:
            db.commit()
            db.refresh(r)
            return r
        except IntegrityError:
            db.rollback()
    return "No se pudo asignar el consecutivo. Intenta de nuevo."


def editar_recibo(db: Session, r: CajaRecibo, user: Empleado, datos: dict) -> str | None:
    """Edita un recibo ACTIVO que todavía no esté firmado. Firmado, solo se puede anular."""
    if r.estado != "ACTIVO":
        return f"Solo se pueden editar recibos ACTIVOS (este está {r.estado})."
    if r.firmado_en:
        return ("Este recibo ya fue firmado por quien lo autorizó y no se puede editar. "
                "Si tiene un error, anúlalo y crea uno nuevo.")
    limpios = _datos_recibo(r.caja, datos)
    if isinstance(limpios, str):
        return limpios
    autoriza, error = _resolver_autoriza(db, r.caja, datos, actual=r, creador_id=r.creado_por_id)
    if error:
        return error
    aviso = _confirmaciones(db, r.caja, limpios, datos, actual=r)
    if aviso:
        return aviso
    for k, v in limpios.items():
        setattr(r, k, v)
    r.autorizado_por_id = autoriza.id if autoriza else None
    _atender_observaciones(r, user)  # la corrección responde a la observación de quien firma
    if not r.numero_factura:
        r.numero_factura = f"{r.caja.prefijo}-{r.consecutivo:04d}"
    r.editado_por_id, r.editado_en = user.id, datetime.utcnow()
    db.commit()
    return None


def anular_recibo(db: Session, r: CajaRecibo, user: Empleado, motivo: str) -> str | None:
    if r.estado == "ANULADO":
        return "Este recibo ya estaba anulado."
    if r.estado == "LEGALIZADO":
        return f"Este recibo está legalizado en el FM{r.fm.numero}. Anula primero ese FM."
    motivo = (motivo or "").strip()
    if len(motivo) < 5:
        return "Escribe el motivo de la anulación (mínimo 5 caracteres)."
    r.estado, r.anulado_por_id, r.anulado_en, r.motivo_anulacion = "ANULADO", user.id, datetime.utcnow(), motivo[:500]
    db.commit()
    return None


def recibos_de_caja(db: Session, caja: CajaMenor, estado: str = "", desde: date | None = None,
                    hasta: date | None = None, limite: int = 2000) -> list[CajaRecibo]:
    q = (db.query(CajaRecibo).options(joinedload(CajaRecibo.caja), joinedload(CajaRecibo.fm),
                                      joinedload(CajaRecibo.creado_por), joinedload(CajaRecibo.anulado_por))
         .filter(CajaRecibo.caja_id == caja.id))
    if estado:
        q = q.filter(CajaRecibo.estado == estado)
    if desde:
        q = q.filter(CajaRecibo.fecha >= desde)
    if hasta:
        q = q.filter(CajaRecibo.fecha <= hasta)
    return q.order_by(CajaRecibo.consecutivo.desc()).limit(limite).all()


def total_pendiente(db: Session, caja: CajaMenor) -> float:
    """Dinero que salió de la caja y aún no vuelve: recibos activos + recibos de FM vigentes sin reembolsar."""
    activos = (db.query(func.coalesce(func.sum(CajaRecibo.valor), 0.0))
               .filter(CajaRecibo.caja_id == caja.id, CajaRecibo.estado == "ACTIVO").scalar() or 0)
    sin_reembolso = (db.query(func.coalesce(func.sum(CajaRecibo.valor), 0.0)).join(CajaFM, CajaRecibo.fm_id == CajaFM.id)
                     .filter(CajaRecibo.caja_id == caja.id, CajaRecibo.estado == "LEGALIZADO", CajaFM.estado == "VIGENTE",
                             CajaFM.reembolsado_en.is_(None)).scalar() or 0)
    return float(activos) + float(sin_reembolso)


def total_sin_reembolsar(db: Session, caja: CajaMenor) -> float:
    return float(db.query(func.coalesce(func.sum(CajaRecibo.valor), 0.0)).join(CajaFM, CajaRecibo.fm_id == CajaFM.id)
                 .filter(CajaRecibo.caja_id == caja.id, CajaRecibo.estado == "LEGALIZADO", CajaFM.estado == "VIGENTE",
                         CajaFM.reembolsado_en.is_(None)).scalar() or 0)


# ---------------- Legalización (FM) ----------------

def serializar_fm(fm: CajaFM, con_recibos: bool = False) -> dict:
    d = {
        "id": fm.id, "numero": fm.numero, "fm": f"FM{fm.numero}", "fecha": fm.fecha_legalizacion.isoformat(),
        "responsable": fm.responsable, "fondo": fm.fondo, "valorEnCaja": fm.valor_en_caja, "ajuste": fm.ajuste,
        "totalPagos": fm.total_pagos, "desglose": json.loads(fm.desglose or "{}"), "estado": fm.estado,
        "cantidad": len(fm.recibos) if fm.estado == "VIGENTE" else None,
        "creadoPor": nombre_propio(fm.creado_por.nombre_completo) if fm.creado_por else fm.creado_por_texto,
        "creadoEn": hora_colombia(fm.creado_en),
        "anuladoPor": nombre_propio(fm.anulado_por.nombre_completo) if fm.anulado_por else "",
        "anuladoEn": hora_colombia(fm.anulado_en), "motivoAnulacion": fm.motivo_anulacion or "",
        "ciudad": fm.caja.ciudad,
        "supervisadoPor": fm.supervisado_por or "", "supervisadoPorId": fm.supervisado_por_id,
        "supervisionEmail": fm.supervision_email or "", "supervisadoEn": hora_colombia(fm.supervisado_en),
        "creadoPorEmail": fm.creado_por.email if fm.creado_por else "",
        "elaboradoPorId": fm.elaborado_por_id, "elaboracionEmail": fm.elaboracion_email or "",
        "elaboradoEn": hora_colombia(fm.elaborado_en),
        "reembolsadoEn": hora_colombia(fm.reembolsado_en), "avisoOk": fm.aviso_ok, "avisoEn": hora_colombia(fm.aviso_en),
        "observaciones": _serializar_observaciones(fm.observaciones)[0],
        "observacionPendiente": _serializar_observaciones(fm.observaciones)[1],
    }
    if con_recibos:
        d["recibos"] = [serializar_recibo(r) for r in fm.recibos]
    return d


def legalizar(db: Session, caja: CajaMenor, user: Empleado, recibo_ids: list[int], responsable: str,
              desglose: dict, ajuste: float, supervisado_por: str = "") -> CajaFM | str:
    if not recibo_ids:
        return "Selecciona al menos un recibo."
    elegido = (supervisado_por or "").strip().upper()
    supervisores = (db.query(Empleado).join(CajaSupervisor, CajaSupervisor.empleado_id == Empleado.id)
                    .filter(CajaSupervisor.caja_id == caja.id, Empleado.activo == 1).all())
    supervisor = next((e for e in supervisores if e.nombre_completo.strip().upper() == elegido), None)
    if supervisores and not supervisor:
        return "Elige quién supervisa el FM (Supervisado por)."
    responsable_emp = responsable_de_caja(db, caja)
    if not responsable_emp:
        return ("La caja no tiene responsable. Un administrador debe asignarlo en Parámetros › Cajas: "
                "es quien firma el «Elaborado por» del FM.")
    if supervisor and supervisor.id == responsable_emp.id:
        return "Quien supervisa el FM debe ser una persona distinta al responsable de la caja (Elaborado por)."
    recibos = db.query(CajaRecibo).filter(CajaRecibo.id.in_(recibo_ids), CajaRecibo.caja_id == caja.id).all()
    if len(recibos) != len(set(recibo_ids)):
        return "Algunos recibos no pertenecen a esta caja."
    no_activos = [f"{caja.prefijo}-{r.consecutivo:04d}" for r in recibos if r.estado != "ACTIVO"]
    if no_activos:
        return f"Estos recibos ya no están ACTIVOS (otra persona los pudo legalizar o anular): {', '.join(no_activos)}"
    sin_firma = [f"{caja.prefijo}-{r.consecutivo:04d}" for r in recibos if r.autorizado_por_id and not r.firmado_en]
    if sin_firma:
        return f"Solo se pueden legalizar recibos firmados por quien los autoriza. Falta la firma de: {', '.join(sin_firma)}"
    desglose = _desglose_valido(desglose or {})
    try:
        ajuste = round(float(ajuste or 0), 2)
    except (TypeError, ValueError):
        ajuste = 0.0
    fm = CajaFM(caja_id=caja.id, numero=caja.fm_siguiente, fecha_legalizacion=(datetime.utcnow() - timedelta(hours=5)).date(),
                responsable=responsable_emp.nombre_completo.strip().upper(), fondo=caja.fondo,
                valor_en_caja=total_desglose(desglose), ajuste=ajuste,
                elaborado_por_id=responsable_emp.id,
                # Si quien legaliza es el responsable, al legalizar queda firmada la elaboración
                elaboracion_email=user.email if user.id == responsable_emp.id else None,
                elaborado_en=datetime.utcnow() if user.id == responsable_emp.id else None,
                total_pagos=round(sum(r.valor for r in recibos) + ajuste, 2),
                desglose=json.dumps(desglose), creado_por_id=user.id,
                supervisado_por=supervisor.nombre_completo.strip().upper() if supervisor else "",
                supervisado_por_id=supervisor.id if supervisor else None)
    db.add(fm)
    db.flush()
    for r in recibos:
        r.estado, r.fm_id = "LEGALIZADO", fm.id
    caja.fm_siguiente = fm.numero + 1
    db.commit()
    db.refresh(fm)
    return fm


def marcar_reembolsado(db: Session, fm: CajaFM, user: Empleado) -> str | None:
    """Tesorería devolvió a la caja el dinero del FM: deja de contar como pendiente en el arqueo."""
    if fm.estado != "VIGENTE":
        return "El FM está anulado."
    if fm.reembolsado_en:
        return "Este FM ya estaba marcado como reembolsado."
    fm.reembolsado_en, fm.reembolsado_por_id = datetime.utcnow(), user.id
    db.commit()
    return None


def anular_fm(db: Session, fm: CajaFM, user: Empleado, motivo: str) -> str | None:
    if fm.estado == "ANULADO":
        return "Este FM ya estaba anulado."
    if fm.reembolsado_en:
        return "Este FM ya fue reembolsado y no se puede anular."
    motivo = (motivo or "").strip()
    if len(motivo) < 5:
        return "Escribe el motivo de la anulación (mínimo 5 caracteres)."
    for r in list(fm.recibos):  # los recibos vuelven a ACTIVO para corregirlos o legalizarlos de nuevo
        r.estado, r.fm_id = "ACTIVO", None
    fm.estado, fm.anulado_por_id, fm.anulado_en, fm.motivo_anulacion = "ANULADO", user.id, datetime.utcnow(), motivo[:500]
    db.commit()
    return None


# ---------------- Arqueo rápido ----------------

def siguiente_arqueo(db: Session, caja: CajaMenor) -> int:
    return (db.query(func.max(CajaArqueo.consecutivo)).filter(CajaArqueo.caja_id == caja.id).scalar() or 0) + 1


def serializar_arqueo(a: CajaArqueo) -> dict:
    return {"id": a.id, "consecutivo": f"ARQ-{a.consecutivo:04d}", "fecha": hora_colombia(a.fecha),
            "responsable": a.responsable, "fondo": a.fondo, "recibosPendientes": a.recibos_pendientes,
            "efectivoContado": a.efectivo_contado, "estado": a.estado, "diferencia": a.diferencia,
            "desglose": json.loads(a.desglose or "{}"),
            "registradoPor": nombre_propio(a.creado_por.nombre_completo) if a.creado_por else a.creado_por_texto}


def crear_arqueo(db: Session, caja: CajaMenor, user: Empleado, responsable: str, desglose: dict) -> CajaArqueo | str:
    responsable_emp = responsable_de_caja(db, caja)
    if not responsable_emp:
        return "La caja no tiene responsable. Un administrador debe asignarlo en Parámetros › Cajas."
    responsable = responsable_emp.nombre_completo
    desglose = _desglose_valido(desglose or {})
    efectivo = total_desglose(desglose)
    pendientes = total_pendiente(db, caja)
    diferencia = round(efectivo + pendientes - caja.fondo, 2)
    estado = "CUADRADO" if diferencia == 0 else ("FALTANTE" if diferencia < 0 else "SOBRANTE")
    for _ in range(3):
        a = CajaArqueo(caja_id=caja.id, consecutivo=siguiente_arqueo(db, caja),
                       responsable=responsable.strip().upper(), fondo=caja.fondo,
                       recibos_pendientes=pendientes, efectivo_contado=efectivo, estado=estado, diferencia=diferencia,
                       desglose=json.dumps(desglose), creado_por_id=user.id)
        db.add(a)
        try:
            db.commit()
            db.refresh(a)
            return a
        except IntegrityError:
            db.rollback()
    return "No se pudo asignar el consecutivo del arqueo. Intenta de nuevo."


# ---------------- Importación desde la hoja de Google (app anterior) ----------------

def _fecha(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto[:10], fmt).date()
        except ValueError:
            continue
    # Formato largo que guardaba la app anterior en los FM: "jueves, 4 de septiembre de 2026"
    import re
    meses = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
             "septiembre": 9, "setiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}
    m = re.search(r"(\d{1,2})\s+de\s+([a-záéíóú]+)\s+de\s+(\d{4})", texto.lower())
    if m and m.group(2) in meses:
        return date(int(m.group(3)), meses[m.group(2)], int(m.group(1)))
    return None


def _fecha_hora(valor) -> datetime | None:
    if isinstance(valor, datetime):
        return valor + timedelta(hours=5)  # hora Colombia -> UTC
    texto = str(valor or "").strip()
    for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y, %H:%M", "%Y-%m-%d %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto[:16].replace(" p. m.", "").replace(" a. m.", ""), fmt) + timedelta(hours=5)
        except ValueError:
            continue
    return None


def _num(valor) -> float:
    try:
        return float(str(valor).replace("$", "").replace(".", "").replace(",", ".").strip()) if isinstance(valor, str) else float(valor or 0)
    except ValueError:
        return 0.0


# ---------------- Firma de quien autoriza ----------------

def necesita_aviso_firma(r: CajaRecibo) -> bool:
    return bool(r.autorizado_por_id and not r.firmado_en and r.estado != "ANULADO")


def notificar_firma_pendiente(recibo_id: int, recordatorio: bool = False) -> None:
    """Aviso por Cliq (solo desde el bot de la empresa) a quien debe autorizar el recibo. Corre en segundo plano.
    Guarda en el recibo si el aviso llegó (aviso_ok) y cuándo (aviso_en)."""
    from . import config
    from .database import SessionLocal
    db = SessionLocal()
    try:
        r = db.get(CajaRecibo, recibo_id)
        if not r or not necesita_aviso_firma(r):
            return
        quien = db.get(Empleado, r.autorizado_por_id)
        valor = f"$ {r.valor:,.0f}".replace(",", ".")
        creador = nombre_propio(r.creado_por.nombre_completo) if r.creado_por else (r.creado_por_texto or "—")
        corregida = next((o for o in reversed(r.observaciones or []) if o.atendida_en and (not r.aviso_en or o.atendida_en > r.aviso_en)), None)
        texto = (f"{'🔔 *Recordatorio* · ' if recordatorio else ''}✍️ *Recibo de caja menor pendiente de tu firma*\n"
                 f"{'✏️ Se corrigió según tu observación: «' + corregida.texto + '»' + chr(10) if corregida else ''}"
                 f"{r.caja.prefijo}-{r.consecutivo:04d} · {valor} · Caja {nombre_propio(r.caja.nombre)}\n"
                 f"Pagado a: {nombre_propio(r.pagado_a)}\n"
                 f"Concepto: {r.concepto}\n"
                 f"Registrado por: {creador}\n"
                 f"Revísalo y fírmalo en: {config.BASE_URL}/caja-menor/{r.caja_id}?tab=firmas")
        r.aviso_ok = 1 if _avisar([quien.email if quien else ""], texto) else 0
        r.aviso_en = datetime.utcnow()
        db.commit()
    except Exception as ex:  # un aviso fallido nunca debe afectar el recibo
        print(f"[Caja menor] Error enviando aviso de firma del recibo #{recibo_id}: {ex}")
    finally:
        db.close()


def _texto_recibo(r: CajaRecibo) -> str:
    valor = f"$ {r.valor:,.0f}".replace(",", ".")
    return (f"{r.caja.prefijo}-{r.consecutivo:04d} · {valor} · Caja {nombre_propio(r.caja.nombre)}\n"
            f"Pagado a: {nombre_propio(r.pagado_a)}\nConcepto: {r.concepto}")


def _avisar(emails: list[str], texto: str) -> bool:
    """Envía por el bot; devuelve True solo si llegó a todos (False si alguno falló o no hay destinatario)."""
    from .zoho_cliq import enviar_cliq
    destinos = list(dict.fromkeys(e for e in emails if e))  # sin repetir
    resultados = [bool(enviar_cliq(email, texto)) for email in destinos]
    return bool(resultados) and all(resultados)


def firmar_elaboracion(db: Session, fm: CajaFM, user: Empleado) -> str | None:
    """Firma "Elaborado por" del responsable de la caja: queda su correo Zoho y la fecha y hora."""
    if fm.elaborado_por_id != user.id:
        return "Solo el responsable de la caja puede firmar la elaboración de este FM."
    if fm.estado != "VIGENTE":
        return "El FM está anulado."
    if fm.elaborado_en:
        return "La elaboración de este FM ya está firmada."
    fm.elaboracion_email, fm.elaborado_en = user.email, datetime.utcnow()
    db.commit()
    return None


def reenviar_aviso_fm(fm_id: int, recordatorio: bool = False) -> None:
    """Vuelve a pedir la firma que le toca al FM: primero la del responsable, luego el visto bueno."""
    from .database import SessionLocal
    db = SessionLocal()
    try:
        fm = db.get(CajaFM, fm_id)
        falta_elaborar = bool(fm and fm.elaborado_por_id and not fm.elaborado_en)
    finally:
        db.close()
    if falta_elaborar:
        notificar_elaboracion_pendiente(fm_id, recordatorio)
    else:
        notificar_supervision_pendiente(fm_id, recordatorio)


HORAS_RECORDATORIO = 24


def enviar_recordatorios() -> int:
    """Recordatorio por Cliq de las firmas pendientes cuyo último aviso tiene más de 24 horas. Devuelve cuántos envió."""
    from .database import SessionLocal
    limite = datetime.utcnow() - timedelta(hours=HORAS_RECORDATORIO)
    db = SessionLocal()
    try:
        recibos = [r.id for r in db.query(CajaRecibo).filter(
            CajaRecibo.autorizado_por_id.isnot(None), CajaRecibo.firmado_en.is_(None), CajaRecibo.estado != "ANULADO",
            func.coalesce(CajaRecibo.aviso_en, CajaRecibo.creado_en) < limite)]
        fms = [f.id for f in db.query(CajaFM).filter(
            CajaFM.estado == "VIGENTE", func.coalesce(CajaFM.aviso_en, CajaFM.creado_en) < limite,
            ((CajaFM.elaborado_por_id.isnot(None)) & (CajaFM.elaborado_en.is_(None)))
            | ((CajaFM.supervisado_por_id.isnot(None)) & (CajaFM.supervisado_en.is_(None))))]
    finally:
        db.close()
    for rid in recibos:
        notificar_firma_pendiente(rid, recordatorio=True)
    for fid in fms:
        reenviar_aviso_fm(fid, recordatorio=True)
    return len(recibos) + len(fms)


def observar_recibo(db: Session, r: CajaRecibo, user: Empleado, texto: str) -> str | None:
    """Quien debe autorizar el recibo no lo firma y deja una observación para que lo corrijan."""
    if r.autorizado_por_id != user.id:
        return "Solo quien debe autorizar este recibo puede dejarle una observación."
    if r.estado == "ANULADO" or r.firmado_en:
        return "El recibo ya no está pendiente de tu firma."
    if r.estado == "LEGALIZADO":
        return "El recibo ya está legalizado; para corregirlo hay que anular su FM."
    texto = (texto or "").strip()
    if len(texto) < 5:
        return "Escribe la observación (mínimo 5 caracteres): qué se debe corregir."
    db.add(CajaObservacion(recibo_id=r.id, autor_id=user.id, texto=texto[:1000]))
    db.commit()
    db.refresh(r)
    return None


def observar_fm(db: Session, fm: CajaFM, user: Empleado, texto: str) -> str | None:
    """Quien supervisa el FM no da el visto bueno y deja una observación para que el responsable lo corrija."""
    if fm.supervisado_por_id != user.id:
        return "Solo quien supervisa este FM puede dejarle una observación."
    if fm.estado != "VIGENTE" or fm.supervisado_en:
        return "El FM ya no está pendiente de tu visto bueno."
    texto = (texto or "").strip()
    if len(texto) < 5:
        return "Escribe la observación (mínimo 5 caracteres): qué se debe corregir."
    db.add(CajaObservacion(fm_id=fm.id, autor_id=user.id, texto=texto[:1000]))
    db.commit()
    db.refresh(fm)
    return None


def corregir_fm(db: Session, fm: CajaFM, user: Empleado, desglose: dict, ajuste: float) -> str | None:
    """El responsable corrige el conteo y el ajuste del FM (antes del visto bueno). Su firma queda renovada.
    Si hay que cambiar los recibos del FM, se anula y se legaliza de nuevo."""
    if fm.elaborado_por_id != user.id:
        return "Solo el responsable de la caja (Elaborado por) puede corregir este FM."
    if fm.estado != "VIGENTE" or fm.supervisado_en or fm.reembolsado_en:
        return "Este FM ya no se puede corregir (está anulado, supervisado o reembolsado)."
    desglose = _desglose_valido(desglose or {})
    try:
        ajuste = round(float(ajuste or 0), 2)
    except (TypeError, ValueError):
        ajuste = 0.0
    fm.desglose, fm.valor_en_caja, fm.ajuste = json.dumps(desglose), total_desglose(desglose), ajuste
    fm.total_pagos = round(sum(r.valor for r in fm.recibos) + ajuste, 2)
    fm.elaboracion_email, fm.elaborado_en = user.email, datetime.utcnow()
    _atender_observaciones(fm, user)
    db.commit()
    return None


def notificar_observacion(recibo_id: int | None = None, fm_id: int | None = None) -> None:
    """Aviso por Cliq (bot) al responsable de la caja y a quien registró/legalizó: hay una observación por corregir."""
    from . import config
    from .database import SessionLocal
    db = SessionLocal()
    try:
        doc = db.get(CajaRecibo, recibo_id) if recibo_id else db.get(CajaFM, fm_id)
        if not doc or not doc.observaciones:
            return
        obs = doc.observaciones[-1]
        responsable = responsable_de_caja(db, doc.caja)
        destinos = [e.email for e in (responsable, doc.creado_por) if e and e.id != obs.autor_id]
        autor = nombre_propio(obs.autor.nombre_completo) if obs.autor else "—"
        if recibo_id:
            valor = f"$ {doc.valor:,.0f}".replace(",", ".")
            cabeza = f"{doc.caja.prefijo}-{doc.consecutivo:04d} · {valor} · Caja {nombre_propio(doc.caja.nombre)}"
            accion = f"Corrígelo (Editar) en: {config.BASE_URL}/caja-menor/{doc.caja_id}"
            titulo = "Observación en recibo de caja menor"
        else:
            cabeza = _texto_fm(doc)
            accion = f"El responsable lo corrige desde Historial FM en: {config.BASE_URL}/caja-menor/{doc.caja_id}?tab=fms"
            titulo = "Observación en formato de reembolso (FM)"
        texto = f"📝 *{titulo}*\n{cabeza}\n{autor} no firmó y dejó esta observación:\n«{obs.texto}»\n{accion}"
        _avisar(destinos, texto)
    except Exception as ex:
        print(f"[Caja menor] Error enviando aviso de observación: {ex}")
    finally:
        db.close()


def dar_visto_bueno(db: Session, fm: CajaFM, user: Empleado) -> str | None:
    """Visto bueno del supervisor del FM: queda su correo Zoho (el de la sesión) y la fecha y hora."""
    if fm.supervisado_por_id != user.id:
        return "Este FM no está asignado a ti para supervisar."
    if fm.estado != "VIGENTE":
        return "El FM está anulado."
    if fm.elaborado_por_id and not fm.elaborado_en:
        return "Primero debe firmar el responsable de la caja (Elaborado por)."
    if fm.supervisado_en:
        return "Este FM ya tiene el visto bueno."
    _atender_observaciones(fm, user)
    fm.supervision_email, fm.supervisado_en = user.email, datetime.utcnow()
    db.commit()
    return None


def _texto_fm(fm: CajaFM) -> str:
    total = f"$ {fm.total_pagos:,.0f}".replace(",", ".")
    return f"FM{fm.numero} · {total} · Caja {nombre_propio(fm.caja.nombre)} · {len(fm.recibos)} recibos"


def notificar_elaboracion_pendiente(fm_id: int, recordatorio: bool = False) -> None:
    """Aviso por Cliq (bot) al responsable de la caja para firmar "Elaborado por" (cuando otra persona legalizó)."""
    from . import config
    from .database import SessionLocal
    db = SessionLocal()
    try:
        fm = db.get(CajaFM, fm_id)
        if not fm or not fm.elaborado_por_id or fm.elaborado_en or fm.estado != "VIGENTE":
            return
        quien = db.get(Empleado, fm.elaborado_por_id)
        creador = nombre_propio(fm.creado_por.nombre_completo) if fm.creado_por else "—"
        texto = (f"{'🔔 *Recordatorio* · ' if recordatorio else ''}📝 *FM pendiente de tu firma como responsable (Elaborado por)*\n"
                 f"{_texto_fm(fm)}\nLegalizado por: {creador}\n"
                 f"Revísalo y fírmalo en: {config.BASE_URL}/caja-menor/{fm.caja_id}?tab=firmas")
        fm.aviso_ok = 1 if _avisar([quien.email if quien else ""], texto) else 0
        fm.aviso_en = datetime.utcnow()
        db.commit()
    except Exception as ex:
        print(f"[Caja menor] Error enviando aviso de elaboración del FM #{fm_id}: {ex}")
    finally:
        db.close()


def notificar_supervision_pendiente(fm_id: int, recordatorio: bool = False) -> None:
    """Aviso por Cliq (bot) a quien debe dar el visto bueno del FM. Corre en segundo plano."""
    from . import config
    from .database import SessionLocal
    db = SessionLocal()
    try:
        fm = db.get(CajaFM, fm_id)
        if not fm or not fm.supervisado_por_id or fm.supervisado_en or fm.estado != "VIGENTE":
            return
        if fm.elaborado_por_id and not fm.elaborado_en:  # se avisa al supervisor cuando el responsable ya firmó
            return
        quien = db.get(Empleado, fm.supervisado_por_id)
        creador = nombre_propio(fm.creado_por.nombre_completo) if fm.creado_por else "—"
        corregida = next((o for o in reversed(fm.observaciones or []) if o.atendida_en and (not fm.aviso_en or o.atendida_en > fm.aviso_en)), None)
        texto = (f"{'🔔 *Recordatorio* · ' if recordatorio else ''}📋 *Formato de reembolso (FM) pendiente de tu visto bueno*\n"
                 f"{'✏️ Se corrigió según tu observación: «' + corregida.texto + '»' + chr(10) if corregida else ''}"
                 f"{_texto_fm(fm)}\nResponsable: {nombre_propio(fm.responsable) or '—'}\nLegalizado por: {creador}\n"
                 f"Revísalo y fírmalo en: {config.BASE_URL}/caja-menor/{fm.caja_id}?tab=firmas")
        fm.aviso_ok = 1 if _avisar([quien.email if quien else ""], texto) else 0
        fm.aviso_en = datetime.utcnow()
        db.commit()
    except Exception as ex:
        print(f"[Caja menor] Error enviando aviso de supervisión del FM #{fm_id}: {ex}")
    finally:
        db.close()


def notificar_supervisado(fm_id: int) -> None:
    """Aviso por Cliq (bot) a quien legalizó el FM: ya tiene el visto bueno."""
    from .database import SessionLocal
    db = SessionLocal()
    try:
        fm = db.get(CajaFM, fm_id)
        if not fm or not fm.supervisado_en or not fm.creado_por or fm.creado_por.email == fm.supervision_email:
            return
        quien = db.get(Empleado, fm.supervisado_por_id) if fm.supervisado_por_id else None
        texto = (f"✅ *FM con visto bueno*\n{_texto_fm(fm)}\n"
                 f"Supervisado por: {nombre_propio(quien.nombre_completo) if quien else fm.supervision_email} "
                 f"({fm.supervision_email}) el {hora_colombia(fm.supervisado_en)}")
        _avisar([fm.creado_por.email], texto)
    except Exception as ex:
        print(f"[Caja menor] Error enviando aviso de FM supervisado #{fm_id}: {ex}")
    finally:
        db.close()


def notificar_fm_anulado(fm_id: int) -> None:
    """Aviso por Cliq (bot) a quien legalizó y a quien supervisa el FM (sin avisar a quien anuló)."""
    from .database import SessionLocal
    db = SessionLocal()
    try:
        fm = db.get(CajaFM, fm_id)
        if not fm or fm.estado != "ANULADO":
            return
        supervisor = db.get(Empleado, fm.supervisado_por_id) if fm.supervisado_por_id else None
        responsable = db.get(Empleado, fm.elaborado_por_id) if fm.elaborado_por_id else None
        anula = fm.anulado_por
        destinos = [e.email for e in (fm.creado_por, responsable, supervisor) if e and (not anula or e.id != anula.id)]
        if not destinos:
            return
        total = f"$ {fm.total_pagos:,.0f}".replace(",", ".")
        texto = (f"🚫 *FM anulado*\nFM{fm.numero} · {total} · Caja {nombre_propio(fm.caja.nombre)}\n"
                 f"Anulado por: {nombre_propio(anula.nombre_completo) if anula else '—'} el {hora_colombia(fm.anulado_en)}\n"
                 f"Motivo: {fm.motivo_anulacion or '—'}\nSus recibos volvieron a quedar activos.")
        _avisar(destinos, texto)
    except Exception as ex:
        print(f"[Caja menor] Error enviando aviso de FM anulado #{fm_id}: {ex}")
    finally:
        db.close()


def notificar_firmado(recibo_id: int) -> None:
    """Aviso por Cliq (bot) a quien registró el recibo: ya lo firmaron. Corre en segundo plano."""
    from .database import SessionLocal
    db = SessionLocal()
    try:
        r = db.get(CajaRecibo, recibo_id)
        if not r or not r.firmado_en or not r.creado_por or r.creado_por.email == r.firma_email:
            return
        quien = db.get(Empleado, r.autorizado_por_id) if r.autorizado_por_id else None
        texto = (f"✅ *Recibo de caja menor firmado*\n{_texto_recibo(r)}\n"
                 f"Firmado por: {nombre_propio(quien.nombre_completo) if quien else r.firma_email} "
                 f"({r.firma_email}) el {hora_colombia(r.firmado_en)}")
        _avisar([r.creado_por.email], texto)
    except Exception as ex:  # un aviso fallido nunca debe afectar el recibo
        print(f"[Caja menor] Error enviando aviso de recibo firmado #{recibo_id}: {ex}")
    finally:
        db.close()


def notificar_anulado(recibo_id: int) -> None:
    """Aviso por Cliq (bot) a quien registró el recibo y a quien lo autoriza (sin avisar a quien anuló)."""
    from .database import SessionLocal
    db = SessionLocal()
    try:
        r = db.get(CajaRecibo, recibo_id)
        if not r or r.estado != "ANULADO":
            return
        autoriza = db.get(Empleado, r.autorizado_por_id) if r.autorizado_por_id else None
        quien_anula = r.anulado_por
        destinatarios = [e.email for e in (r.creado_por, autoriza) if e and (not quien_anula or e.id != quien_anula.id)]
        if not destinatarios:
            return
        texto = (f"🚫 *Recibo de caja menor anulado*\n{_texto_recibo(r)}\n"
                 f"{'Estaba firmado por ' + r.firma_email + chr(10) if r.firmado_en else ''}"
                 f"Anulado por: {nombre_propio(quien_anula.nombre_completo) if quien_anula else '—'} el {hora_colombia(r.anulado_en)}\n"
                 f"Motivo: {r.motivo_anulacion or '—'}")
        _avisar(destinatarios, texto)
    except Exception as ex:
        print(f"[Caja menor] Error enviando aviso de recibo anulado #{recibo_id}: {ex}")
    finally:
        db.close()


def firmar_recibo(db: Session, r: CajaRecibo, user: Empleado) -> str | None:
    """Firma del autorizador: queda su correo Zoho (con el que inició sesión) y la fecha y hora."""
    if r.autorizado_por_id != user.id:
        return "Este recibo no está asignado a ti para firmar."
    if r.estado == "ANULADO":
        return "El recibo está anulado."
    if r.firmado_en:
        return "Este recibo ya está firmado."
    _atender_observaciones(r, user)
    r.firma_email, r.firmado_en = user.email, datetime.utcnow()
    db.commit()
    return None


def importar_hoja_google(db: Session, contenido: bytes, user: Empleado) -> dict:
    """Importa el Excel descargado de la hoja de Google de la app anterior (Archivo > Descargar > .xlsx):
    hojas "REGISTRO <CAJA>" (recibos), "ARQUEOS_<CAJA>" (FMs) y "REGISTRO_ARQUEOS_<CAJA>" (arqueos rápidos).
    No duplica: omite recibos/FMs/arqueos cuyo consecutivo ya exista en la caja."""
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(contenido), data_only=True, read_only=True)
    empleados = {e.email.lower(): e for e in db.query(Empleado).all() if e.email}
    resumen = {"recibos": 0, "fms": 0, "fmsAnulados": 0, "arqueos": 0, "omitidos": 0, "cajasCreadas": []}

    def caja_por_nombre(nombre: str) -> CajaMenor:
        nombre = nombre.strip().upper()
        caja = db.query(CajaMenor).filter(CajaMenor.nombre == nombre).first()
        if not caja:
            caja = CajaMenor(nombre=nombre, prefijo=nombre[:4], orden=db.query(CajaMenor).count() + 1)
            db.add(caja)
            db.flush()
            resumen["cajasCreadas"].append(nombre)
        return caja

    def autor(texto) -> tuple[int | None, str]:
        texto = str(texto or "").strip()
        correo = texto.split(":")[-1].strip().lower()
        emp = empleados.get(correo)
        return (emp.id if emp else None), texto

    hojas = {ws.title.strip().upper(): ws for ws in wb.worksheets}
    fms_por_caja: dict[tuple[int, int], CajaFM] = {}

    # 1) FMs (hojas ARQUEOS_<CAJA>, sin "REGISTRO_")
    for titulo, ws in hojas.items():
        if not titulo.startswith("ARQUEOS_"):
            continue
        caja = caja_por_nombre(titulo.replace("ARQUEOS_", ""))
        for fila in list(ws.iter_rows(values_only=True))[1:]:
            if not fila or not fila[1]:
                continue
            try:
                numero = int(str(fila[1]).upper().replace("FM", "").strip())
            except ValueError:
                continue
            if db.query(CajaFM).filter(CajaFM.caja_id == caja.id, CajaFM.numero == numero).first():
                resumen["omitidos"] += 1
                continue
            valores = list(fila) + [None] * 18
            desglose = {str(d): int(_num(valores[5 + i])) for i, d in enumerate(MONEDAS + BILLETES)}
            creador_id, creador_txt = autor(valores[17])
            fm = CajaFM(caja_id=caja.id, numero=numero, fecha_legalizacion=_fecha(valores[0]) or date.today(),
                        responsable=str(valores[2] or "").upper(), fondo=caja.fondo, valor_en_caja=_num(valores[3]),
                        reembolsado_en=datetime.utcnow(),  # la app anterior no llevaba el reembolso: se asumen reembolsados
                        total_pagos=_num(valores[4]), ajuste=_num(valores[16]), desglose=json.dumps(desglose),
                        creado_por_id=creador_id, creado_por_texto=creador_txt)
            db.add(fm)
            db.flush()
            fms_por_caja[(caja.id, numero)] = fm
            resumen["fms"] += 1

    # 2) Recibos (hojas con el nombre de la caja, que no sean de arqueos)
    for titulo, ws in hojas.items():
        if "ARQUEO" in titulo or titulo == "INF":
            continue
        filas = list(ws.iter_rows(values_only=True))
        if not filas or "CONS" not in [str(c or "").strip().upper() for c in filas[0]]:
            continue
        for fila in filas[1:]:
            valores = list(fila or []) + [None] * 16
            if valores[5] in (None, ""):
                continue
            try:
                consecutivo = int(valores[5])
            except (TypeError, ValueError):
                continue
            caja = caja_por_nombre(str(valores[1] or titulo.replace("REGISTRO", "")))
            if valores[4]:
                caja.prefijo = str(valores[4]).strip().upper()
            if db.query(CajaRecibo).filter(CajaRecibo.caja_id == caja.id, CajaRecibo.consecutivo == consecutivo).first():
                resumen["omitidos"] += 1
                continue
            estado = str(valores[12] or "ACTIVO").strip().upper()
            estado = estado if estado in ("ACTIVO", "ANULADO", "LEGALIZADO") else "ACTIVO"
            fm = None
            if estado == "LEGALIZADO" and valores[13]:
                try:
                    fm = fms_por_caja.get((caja.id, int(str(valores[13]).upper().replace("FM", "")))) or \
                        db.query(CajaFM).filter(CajaFM.caja_id == caja.id,
                                                CajaFM.numero == int(str(valores[13]).upper().replace("FM", ""))).first()
                except ValueError:
                    fm = None
            creador_id, creador_txt = autor(valores[15])
            db.add(CajaRecibo(
                caja_id=caja.id, consecutivo=consecutivo, ciudad=str(valores[2] or caja.ciudad).upper(),
                fecha=_fecha(valores[3]) or date.today(), identificacion=str(valores[6] or ""),
                pagado_a=str(valores[7] or ""), valor=_num(valores[8]), valor_letras=numero_a_letras(_num(valores[8])),
                concepto=str(valores[10] or ""), numero_factura=str(valores[11] or ""), estado=estado,
                fm_id=fm.id if fm else None, creado_por_id=creador_id, creado_por_texto=creador_txt,
                creado_en=_fecha_hora(valores[0]) or datetime.utcnow(),
                motivo_anulacion="Anulado en la app anterior (Google)" if estado == "ANULADO" else None))
            resumen["recibos"] += 1
    db.flush()

    # 3) FMs sin recibos asociados = FMs que se anularon en la app anterior
    for fm in fms_por_caja.values():
        if not fm.recibos:
            fm.estado, fm.motivo_anulacion = "ANULADO", "Anulado en la app anterior (Google)"
            resumen["fmsAnulados"] += 1

    # 4) Arqueos rápidos (hojas REGISTRO_ARQUEOS_<CAJA>)
    for titulo, ws in hojas.items():
        if not titulo.startswith("REGISTRO_ARQUEOS_"):
            continue
        caja = caja_por_nombre(titulo.replace("REGISTRO_ARQUEOS_", ""))
        for fila in list(ws.iter_rows(values_only=True))[1:]:
            valores = list(fila or []) + [None] * 11
            try:
                consecutivo = int(str(valores[1] or "").upper().replace("ARQ-", ""))
            except ValueError:
                continue
            if db.query(CajaArqueo).filter(CajaArqueo.caja_id == caja.id, CajaArqueo.consecutivo == consecutivo).first():
                resumen["omitidos"] += 1
                continue
            try:
                desglose = _desglose_valido(json.loads(valores[9] or "{}"))
            except (TypeError, ValueError):
                desglose = _desglose_valido({})
            creador_id, creador_txt = autor(valores[10])
            db.add(CajaArqueo(caja_id=caja.id, consecutivo=consecutivo, fecha=_fecha_hora(valores[0]) or datetime.utcnow(),
                              responsable=str(valores[3] or "").upper(), fondo=_num(valores[4]),
                              recibos_pendientes=_num(valores[5]), efectivo_contado=_num(valores[6]),
                              estado=str(valores[7] or "CUADRADO").upper(), diferencia=_num(valores[8]),
                              desglose=json.dumps(desglose), creado_por_id=creador_id, creado_por_texto=creador_txt))
            resumen["arqueos"] += 1

    # 5) Próximo FM de cada caja = mayor importado + 1
    for caja in db.query(CajaMenor).all():
        maximo = db.query(func.max(CajaFM.numero)).filter(CajaFM.caja_id == caja.id).scalar()
        if maximo and caja.fm_siguiente <= maximo:
            caja.fm_siguiente = maximo + 1
    db.commit()
    return resumen
