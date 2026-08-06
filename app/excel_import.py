"""Importación y exportación de empleados/reportes en Excel."""
from datetime import datetime, date
from io import BytesIO
from openpyxl import load_workbook, Workbook
from sqlalchemy.orm import Session
from .models import Empleado, Solicitud, MODULOS_VALIDOS

COLUMNAS = ["nombres", "apellidos", "fecha_nacimiento", "fecha_inicio_empresa", "empresa", "cargo", "area",
            "identificacion", "correo", "rol", "num_aprobaciones", "dias_vacaciones", "modulos", "activo",
            "aprobador1_correo", "aprobador2_correo"]

REQUERIDAS = ["nombres", "apellidos", "empresa", "cargo", "area", "identificacion", "correo"]
ROLES_VALIDOS = ("empleado", "aprobador", "admin")
ACTIVOS_TXT = {"1": 1, "si": 1, "sí": 1, "activo": 1, "0": 0, "no": 0, "inactivo": 0}


def _fecha(valor) -> date | None:
    if valor in (None, ""):
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(valor).strip(), fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Fecha inválida: {valor!r} (usa AAAA-MM-DD o DD/MM/AAAA)")


def importar_empleados(db: Session, contenido: bytes) -> dict:
    """Importa/actualiza empleados desde un xlsx. Upsert por identificación."""
    wb = load_workbook(BytesIO(contenido), data_only=True)
    ws = wb.active
    encabezados = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    faltantes = [c for c in REQUERIDAS if c not in encabezados]
    if faltantes:
        return {"creados": 0, "actualizados": 0,
                "errores": [f"Faltan columnas obligatorias: {', '.join(faltantes)}. "
                            f"Descarga la plantilla oficial."]}
    idx = {c: encabezados.index(c) for c in COLUMNAS if c in encabezados}

    creados, actualizados, errores = 0, 0, []
    pendientes_aprobadores = []  # (identificacion, correo_apr1, correo_apr2)

    for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(fila):
            continue
        def val(col):
            i = idx.get(col)
            v = fila[i] if i is not None and i < len(fila) else None
            return str(v).strip() if v not in (None, "") else ""
        try:
            identificacion = val("identificacion")
            correo = val("correo").lower()
            if not identificacion or not correo:
                raise ValueError("identificación y correo son obligatorios")
            num_apr = int(float(val("num_aprobaciones") or "1"))
            if num_apr not in (1, 2):
                raise ValueError("num_aprobaciones debe ser 1 o 2")

            emp = db.query(Empleado).filter(Empleado.identificacion == identificacion).first()
            nuevo = emp is None
            if nuevo:
                emp = Empleado(identificacion=identificacion)
                db.add(emp)
            emp.nombres = val("nombres")
            emp.apellidos = val("apellidos")
            i_fecha = idx.get("fecha_nacimiento")
            emp.fecha_nacimiento = _fecha(fila[i_fecha] if i_fecha is not None and i_fecha < len(fila) else None)
            i_fecha_ini = idx.get("fecha_inicio_empresa")
            emp.fecha_inicio_empresa = _fecha(
                fila[i_fecha_ini] if i_fecha_ini is not None and i_fecha_ini < len(fila) else None)
            emp.empresa = val("empresa")
            emp.cargo = val("cargo")
            emp.area = val("area")
            emp.email = correo
            emp.num_aprobaciones = num_apr

            rol_txt = val("rol").lower()
            if rol_txt:
                if rol_txt not in ROLES_VALIDOS:
                    raise ValueError(f"rol inválido: '{rol_txt}' (usa empleado, aprobador o admin)")
                emp.rol = rol_txt
            elif nuevo:
                emp.rol = "empleado"

            dias_vac_txt = val("dias_vacaciones")
            if dias_vac_txt:
                emp.dias_vacaciones = float(dias_vac_txt)
            elif nuevo:
                emp.dias_vacaciones = 0

            modulos_txt = val("modulos")
            if modulos_txt:
                slugs = [s.strip().lower() for s in modulos_txt.replace(";", ",").split(",") if s.strip()]
                invalidos = [s for s in slugs if s not in MODULOS_VALIDOS]
                if invalidos:
                    raise ValueError(f"módulo inválido: {', '.join(invalidos)} (usa people, compras o sst)")
                emp.modulos = ",".join(slugs)
            elif nuevo:
                emp.modulos = "people"

            activo_txt = val("activo").lower()
            if activo_txt:
                if activo_txt not in ACTIVOS_TXT:
                    raise ValueError(f"activo inválido: '{activo_txt}' (usa si/no o 1/0)")
                emp.activo = ACTIVOS_TXT[activo_txt]
            elif nuevo:
                emp.activo = 1

            db.flush()
            pendientes_aprobadores.append(
                (identificacion, val("aprobador1_correo").lower(), val("aprobador2_correo").lower()))
            creados += 1 if nuevo else 0
            actualizados += 0 if nuevo else 1
        except Exception as e:
            errores.append(f"Fila {n}: {e}")

    # Segundo paso: vincular aprobadores (pueden venir en el mismo archivo)
    for identificacion, c1, c2 in pendientes_aprobadores:
        emp = db.query(Empleado).filter(Empleado.identificacion == identificacion).first()
        for correo_apr, attr in ((c1, "aprobador1_id"), (c2, "aprobador2_id")):
            if not correo_apr:
                continue
            apr = db.query(Empleado).filter(Empleado.email == correo_apr).first()
            if apr:
                setattr(emp, attr, apr.id)
            else:
                errores.append(f"{emp.identificacion}: aprobador '{correo_apr}' no existe en el sistema")

    db.commit()
    return {"creados": creados, "actualizados": actualizados, "errores": errores}


def generar_plantilla() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Empleados"
    ws.append(COLUMNAS)
    ws.append(["Ana María", "Pérez Gómez", "1990-05-14", "2022-03-01", "Nuvia Smiles Colombia SAS", "Coordinadora",
               "Admin", "1032456789", "ana.perez@empresa.com", "empleado", 1, 0, "people", "si",
               "jefe@empresa.com", ""])
    ws.append(["Carlos", "Ruiz López", "1985-11-02", "2020-07-15", "Nuvia Design Colombia SAS", "Gerente",
               "Operativa", "79845123", "carlos.ruiz@empresa.com", "aprobador", 2, 15, "people", "si",
               "jefe@empresa.com", "direccion@empresa.com"])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar_reporte(db: Session, desde: date | None, hasta: date | None,
                     area: str = "", empresa: str = "", estado: str = "") -> bytes:
    q = db.query(Solicitud).join(Empleado, Solicitud.empleado_id == Empleado.id)
    if desde:
        q = q.filter(Solicitud.fecha_inicio >= desde)
    if hasta:
        q = q.filter(Solicitud.fecha_inicio <= hasta)
    if area:
        q = q.filter(Empleado.area == area)
    if empresa:
        q = q.filter(Empleado.empresa == empresa)
    if estado:
        q = q.filter(Solicitud.estado == estado)

    wb = Workbook()
    ws = wb.active
    ws.title = "Permisos"
    ws.append(["ID", "Empleado", "Identificación", "Empresa", "Área", "Cargo", "Tipo de permiso",
               "Fecha inicio", "Fecha fin", "Días", "Estado", "Motivo", "Creada",
               "Aprobador 1", "Decisión 1", "Aprobador 2", "Decisión 2"])
    for s in q.order_by(Solicitud.fecha_inicio.desc()).all():
        a1 = next((a for a in s.aprobaciones if a.nivel == 1), None)
        a2 = next((a for a in s.aprobaciones if a.nivel == 2), None)
        ws.append([s.id, s.empleado.nombre_completo, s.empleado.identificacion, s.empleado.empresa,
                   s.empleado.area, s.empleado.cargo, s.tipo.nombre, s.fecha_inicio, s.fecha_fin,
                   s.dias, s.estado_texto, s.motivo, s.creada_en.strftime("%Y-%m-%d %H:%M"),
                   a1.aprobador.nombre_completo if a1 else "", a1.decision if a1 else "",
                   a2.aprobador.nombre_completo if a2 else "", a2.decision if a2 else ""])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
