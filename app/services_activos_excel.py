"""Módulo Activos › Design: exportar el inventario a Excel y comparar/sincronizar contra un libro
subido. Usa openpyxl igual que excel_import.py (People) -- mismo patrón de la plataforma, no el
estilo del proyecto original (pandas)."""
from io import BytesIO
from datetime import date, datetime
from openpyxl import Workbook
from sqlalchemy.orm import Session
from .models import Empleado
from .models_activos import ActivosEquipo, ActivosDongle
from . import services_activos_consultas as consultas
from . import services_activos as sa

HOJAS_EQUIPOS = ["Asset Tag", "Tipo", "Marca", "Modelo", "Serial", "Sistema operativo", "Usuario", "Área",
                 "Departamento", "Estado", "Dongle", "Situación del dongle"]


def _hoja_equipos(wb: Workbook, filas: list[dict]) -> None:
    ws = wb.active
    ws.title = "Equipos"
    ws.append(HOJAS_EQUIPOS)
    for f in filas:
        ws.append([f["asset_tag"], f["tipo_equipo"], f["marca"], f["modelo"], f["serial_number"], f["sistema_operativo"],
                  f["usuario"], f["area"], f["departamento"], f["estado"], f["dongle"], f["situacion_dongle"]])


def exportar_equipos(db: Session, busqueda: str = "", area: str | None = None, departamento: str | None = None,
                     estado: str | None = None, tipo: str | None = None, dongle: str | None = None) -> bytes:
    filas = consultas.listar_equipos(db, busqueda, area, departamento, estado, tipo, dongle)
    wb = Workbook()
    _hoja_equipos(wb, filas)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar_inventario(db: Session) -> bytes:
    wb = Workbook()
    _hoja_equipos(wb, consultas.listar_equipos(db))

    ws = wb.create_sheet("Usuarios")
    ws.append(["Usuario", "JumpCloud", "Cargo", "Área", "Departamento", "Grupo (líder)", "Equipos asignados"])
    for u in consultas.listar_usuarios(db):
        ws.append([u["nombre_completo"], u["jumpcloud_username"], u["cargo"], u["area"], u["departamento"],
                  u["lider"], u["equipos_asignados"]])

    ws = wb.create_sheet("Dongles")
    ws.append(["Número", "Serial físico", "Versión", "Módulos", "Equipo asignado", "Estado"])
    for d in consultas.listar_dongles(db):
        ws.append([d["numero_dongle"], d["serial_number"], d["version_exocad_compatible"], d["modulos"], d["equipo"], d["estado_dongle"]])

    ws = wb.create_sheet("Mantenimientos")
    ws.append(["Equipo", "Usuario", "Tipo", "Programado", "Realizado", "Estado", "Descripción", "Responsable"])
    for m in consultas.listar_mantenimientos(db):
        ws.append([m["asset_tag"], m["usuario"], m["tipo_mantenimiento"], m["fecha_programada"], m["fecha_realizada"],
                  m["estado_mantenimiento"], m["descripcion_falla"], m["responsable"]])

    ws = wb.create_sheet("Retirados")
    ws.append(["Asset Tag", "Tipo", "Marca", "Modelo", "Serial", "Fecha de retiro", "Motivo"])
    for e in consultas.listar_equipos_retirados(db):
        ws.append([e["asset_tag"], e["tipo_equipo"], e["marca"], e["modelo"], e["serial_number"], e["fecha_retiro"], e["motivo_retiro"]])

    ws = wb.create_sheet("Usuarios retirados")
    ws.append(["Usuario", "Cargo", "Área", "Departamento", "Fecha de retiro", "Motivo"])
    for u in consultas.listar_usuarios_retirados(db):
        ws.append([u["nombre_completo"], u["cargo"], u["area"], u["departamento"], u["fecha_retiro"], u["motivo_retiro"]])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


CAMPOS_EQUIPO_EXCEL = {"Marca": "marca", "Modelo": "modelo_referencia", "Serial": "serial_number",
                       "Sistema operativo": "sistema_operativo"}


def comparar_inventario_excel(db: Session, contenido: bytes) -> dict:
    """Compara, por Asset Tag, los equipos del libro subido contra los del inventario (campos
    editables: marca, modelo, serial, sistema operativo). Devuelve las diferencias encontradas; no
    cambia nada hasta que se llame a sincronizar_cambios con lo que el usuario elija aplicar."""
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(contenido), data_only=True)
    ws = wb.active
    encabezados = [str(c.value).strip() if c.value else "" for c in ws[1]]
    idx = {h: i for i, h in enumerate(encabezados)}
    if "Asset Tag" not in idx:
        raise ValueError("El libro debe tener una columna 'Asset Tag'.")
    actuales = {e.asset_tag: e for e in db.query(ActivosEquipo).filter(ActivosEquipo.retirado.is_(False)).all()}
    diferencias = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        tag = str(fila[idx["Asset Tag"]] or "").strip()
        if not tag or tag not in actuales:
            continue
        equipo = actuales[tag]
        for columna, campo in CAMPOS_EQUIPO_EXCEL.items():
            if columna not in idx:
                continue
            valor_excel = str(fila[idx[columna]] or "").strip()
            valor_actual = str(getattr(equipo, campo) or "").strip()
            if valor_excel and valor_excel != valor_actual:
                diferencias.append({"tipo": "EQUIPO", "identificador": tag, "campo": columna,
                                   "actual": valor_actual, "excel": valor_excel, "accion": "ACTUALIZAR"})
    return {"diferencias": diferencias, "total": len(diferencias)}


def sincronizar_cambios(db: Session, actor: Empleado, cambios: list[dict]) -> int:
    """Aplica los cambios que el usuario seleccionó en la comparación (mismo formato que devuelve
    comparar_inventario_excel): por ahora, solo equipos (marca/modelo/serial/sistema operativo)."""
    aplicados = 0
    campo_por_columna = {v2: v1 for v1, v2 in {"Marca": "marca", "Modelo": "modelo_referencia", "Serial": "serial_number",
                                               "Sistema operativo": "sistema_operativo"}.items()}
    equipos_tocados: dict[str, dict] = {}
    for c in cambios:
        if c.get("tipo") != "EQUIPO":
            continue
        tag = c.get("identificador")
        campo = CAMPOS_EQUIPO_EXCEL.get(c.get("campo"))
        if not tag or not campo:
            continue
        equipos_tocados.setdefault(tag, {})[campo] = c.get("excel")
    for tag, cambios_equipo in equipos_tocados.items():
        e = db.get(ActivosEquipo, tag)
        if not e or e.retirado:
            continue
        sa.actualizar_equipo(db, actor, tag, cambios_equipo.get("tipo_equipo", e.tipo_equipo),
                             cambios_equipo.get("marca", e.marca) or e.marca,
                             cambios_equipo.get("modelo_referencia", e.modelo_referencia) or e.modelo_referencia,
                             cambios_equipo.get("serial_number", e.serial_number) or e.serial_number,
                             cambios_equipo.get("sistema_operativo", e.sistema_operativo),
                             e.observaciones, e.procesador, e.memoria_ram, e.almacenamiento, e.aplica_dongle)
        aplicados += 1
    return aplicados
