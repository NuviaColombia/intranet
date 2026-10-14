"""Cambios digitales (Dashboard › Cambios digitales, 13-oct-2026): registro de los cambios que pide el doctor o el laboratorio
sobre un caso (REPORTE CAMBIOS DIGITAL.xlsx). Se puede buscar, filtrar, ordenar, editar en la celda, agregar, quitar,
importar desde el Excel y descargar. No depende del periodo del Dashboard: es un registro."""
import csv
import io
import re
import unicodedata
from datetime import date, datetime

from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from .models_design import DesignCambioDigital

CAMPOS = ("fecha", "orden", "paciente", "centro", "doctor", "producto", "tipo_cambio", "detalle", "origen")
ORIGENES = ("DOC", "LAB")
ORDENABLES = {"fecha": DesignCambioDigital.fecha, "orden": DesignCambioDigital.orden, "paciente": DesignCambioDigital.paciente,
              "centro": DesignCambioDigital.centro, "doctor": DesignCambioDigital.doctor, "producto": DesignCambioDigital.producto,
              "tipo_cambio": DesignCambioDigital.tipo_cambio, "origen": DesignCambioDigital.origen}
TOPE_TEXTO = {"orden": 40, "paciente": 200, "centro": 120, "doctor": 200, "producto": 200, "tipo_cambio": 120, "detalle": 4000}


def _sin_acentos(t: str) -> str:
    t = unicodedata.normalize("NFD", t or "")
    return "".join(c for c in t if unicodedata.category(c) != "Mn").lower()


def _texto(v, campo: str) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = " ".join(str(v).split())
    return t[:TOPE_TEXTO.get(campo, 200)]


def normalizar(campo: str, valor):
    """Valida y limpia un valor de cualquier campo. Lanza ValueError con un mensaje para la persona."""
    if campo not in CAMPOS:
        raise ValueError("Ese campo no se puede editar.")
    if campo == "fecha":
        if valor in (None, ""):
            return None
        if isinstance(valor, datetime):
            return valor.date()
        if isinstance(valor, date):
            return valor
        t = str(valor).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%m-%d-%Y", "%m/%d/%Y"):   # el Excel trae ambos órdenes
            try:
                return datetime.strptime(t, fmt).date()
            except ValueError:
                pass
        raise ValueError("La fecha no es válida (usa dd/mm/aaaa).")
    if campo == "origen":
        t = _texto(valor, campo).upper()
        if t and t not in ORIGENES:
            raise ValueError("El origen debe ser DOC (doctor) o LAB (laboratorio).")
        return t
    return _texto(valor, campo)


_TRADUCCIONES: dict | None = None


def traduccion_detalle(texto: str) -> str:
    """Traducción al español de un detalle escrito en inglés (seed_data/cambios_detalle_es.json). "" si no hay."""
    global _TRADUCCIONES
    if _TRADUCCIONES is None:
        import json
        from pathlib import Path
        ruta = Path(__file__).resolve().parent / "seed_data" / "cambios_detalle_es.json"
        _TRADUCCIONES = {" ".join(k.split()): v for k, v in json.loads(ruta.read_text(encoding="utf-8")).items()} if ruta.exists() else {}
    return _TRADUCCIONES.get(" ".join((texto or "").split()), "")


def _fila_a_dict(r: DesignCambioDigital) -> dict:
    return {"id": r.id, "fecha": r.fecha.isoformat() if r.fecha else "", "orden": r.orden or "", "paciente": r.paciente or "",
            "centro": r.centro or "", "doctor": r.doctor or "", "producto": r.producto or "", "tipo_cambio": r.tipo_cambio or "",
            "detalle": r.detalle or "", "detalleEs": traduccion_detalle(r.detalle), "origen": r.origen or ""}


def listar(db: Session, q: str = "", desde: date | None = None, hasta: date | None = None, origen: str = "",
           tipo: str = "", centro: str = "", orden: str = "fecha", dir_: str = "desc", limite: int = 50, offset: int = 0) -> dict:
    query = db.query(DesignCambioDigital)
    if desde:
        query = query.filter(DesignCambioDigital.fecha >= desde)
    if hasta:
        query = query.filter(DesignCambioDigital.fecha <= hasta)
    if origen in ORIGENES:
        query = query.filter(DesignCambioDigital.origen == origen)
    if tipo:
        query = query.filter(DesignCambioDigital.tipo_cambio == tipo)
    if centro:
        query = query.filter(DesignCambioDigital.centro == centro)
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(or_(DesignCambioDigital.orden.ilike(like), DesignCambioDigital.paciente.ilike(like),
                                 DesignCambioDigital.doctor.ilike(like), DesignCambioDigital.producto.ilike(like),
                                 DesignCambioDigital.detalle.ilike(like), DesignCambioDigital.centro.ilike(like)))
    total = query.count()
    col = ORDENABLES.get(orden, DesignCambioDigital.fecha)
    orden_sql = col.asc().nullslast() if dir_ == "asc" else col.desc().nullslast()
    filas = query.order_by(orden_sql, DesignCambioDigital.id.desc()).offset(max(0, offset)).limit(max(1, min(limite, 200))).all()
    base = db.query(DesignCambioDigital)
    kpis = {"total": base.count(), "doc": base.filter(DesignCambioDigital.origen == "DOC").count(),
            "lab": base.filter(DesignCambioDigital.origen == "LAB").count(),
            "centros": db.query(func.count(func.distinct(DesignCambioDigital.centro))).filter(DesignCambioDigital.centro != "").scalar() or 0}
    tipos = [t for (t,) in db.query(DesignCambioDigital.tipo_cambio).filter(DesignCambioDigital.tipo_cambio != "").distinct().order_by(DesignCambioDigital.tipo_cambio)]
    centros = [c for (c,) in db.query(DesignCambioDigital.centro).filter(DesignCambioDigital.centro != "").distinct().order_by(DesignCambioDigital.centro)]
    return {"total": total, "items": [_fila_a_dict(r) for r in filas], "kpis": kpis, "tipos": tipos, "centros": centros}


def crear(db: Session, datos: dict, por: str) -> DesignCambioDigital:
    r = DesignCambioDigital(creado_por=(por or "")[:150], creado_en=datetime.utcnow(), actualizado_en=datetime.utcnow())
    for c in CAMPOS:
        setattr(r, c, normalizar(c, datos.get(c)))
    if not (r.paciente or r.orden):
        raise ValueError("Escribe al menos el número de orden o el paciente.")
    if r.origen is None:
        r.origen = ""
    db.add(r)
    db.commit()
    return r


def actualizar(db: Session, id_: int, campo: str, valor) -> DesignCambioDigital | None:
    r = db.get(DesignCambioDigital, id_)
    if not r:
        return None
    setattr(r, campo, normalizar(campo, valor))
    r.actualizado_en = datetime.utcnow()
    db.commit()
    return r


def _clave_centro(t: str) -> str:
    t = unicodedata.normalize("NFD", t or "")
    t = "".join(c for c in t if unicodedata.category(c) != "Mn").lower()
    return re.sub(r"[^a-z0-9]", "", t)


def doctores_de_centro(db: Session, centro: str) -> list[str]:
    """Doctores asignados a ESE centro en Pre-Approved. Cada centro de una hoja agrupa los `span` doctores que siguen
    (igual que la pantalla de Pre-Approved y la ventana del Schedule); no se mezclan los de otros centros."""
    from .models_design import DesignPreApprovedSheet
    k = _clave_centro(centro)
    if not k:
        return []
    nombres, vistos = [], set()
    for h in db.query(DesignPreApprovedSheet).all():
        docs, i = h.doctores, 0
        for c in h.centros:
            span = max(1, c.span or 1)
            grupo = docs[i:i + span]
            i += span
            if _clave_centro(c.nombre) != k:
                continue
            for d in grupo:
                n = " ".join((d.nombre or "").split())
                if n and n.lower() not in vistos:
                    vistos.add(n.lower())
                    nombres.append(n)
    return nombres


def buscar_orden(db: Session, numero: str, equipos: set[int] | None) -> dict:
    """Busca la orden en el Schedule por su número (solo en los equipos que la persona ve; admins: todos).
    Si hay varias con ese número, toma la más reciente y avisa cuántas hay."""
    from .models_design import DesignOrden, DesignTeam
    n = " ".join((numero or "").split()).upper()
    if not n:
        return {"encontrada": False, "motivo": "Escribe el número de orden."}
    q = (db.query(DesignOrden, DesignTeam).join(DesignTeam, DesignOrden.team_id == DesignTeam.id)
         .filter(func.upper(func.trim(DesignOrden.orden)) == n))
    if equipos is not None:
        q = q.filter(DesignOrden.team_id.in_(equipos or [-1]))
    filas = q.order_by(DesignOrden.fecha.desc(), DesignOrden.id.desc()).all()
    if not filas:
        return {"encontrada": False, "motivo": f"No encontré la orden {n} en los schedules que puedes ver."}
    o, t = filas[0]
    return {"encontrada": True, "orden": n, "fecha": o.fecha.isoformat() if o.fecha else "", "paciente": o.paciente or "",
            "centro": o.centro or "", "producto": o.producto or "", "equipo": t.nombre,
            "coincidencias": len(filas)}


REQUERIDOS_FORM = ("doctor", "tipo_cambio", "detalle", "origen")


def crear_desde_orden(db: Session, numero: str, datos: dict, equipos: set[int] | None, por: str) -> DesignCambioDigital:
    """Guarda un cambio tomando los datos de la orden del Schedule. Todos los campos del formulario son obligatorios."""
    from .models_design import DesignOrden, DesignTeam
    for c in REQUERIDOS_FORM:
        if not str(datos.get(c) or "").strip():
            raise ValueError("Falta completar: " + {"doctor": "el doctor", "tipo_cambio": "el tipo de cambio",
                                                   "detalle": "el detalle", "origen": "el origen (DOC o LAB)"}[c] + ".")
    b = buscar_orden(db, numero, equipos)
    if not b["encontrada"]:
        raise ValueError(b["motivo"])
    q = (db.query(DesignOrden).join(DesignTeam, DesignOrden.team_id == DesignTeam.id)
         .filter(func.upper(func.trim(DesignOrden.orden)) == b["orden"]))
    if equipos is not None:
        q = q.filter(DesignOrden.team_id.in_(equipos or [-1]))
    o = q.order_by(DesignOrden.fecha.desc(), DesignOrden.id.desc()).first()
    fila = {"fecha": o.fecha, "orden": b["orden"], "paciente": o.paciente or "", "centro": o.centro or "",
            "producto": o.producto or "", "doctor": datos.get("doctor", ""), "tipo_cambio": datos.get("tipo_cambio", ""),
            "detalle": datos.get("detalle", ""), "origen": datos.get("origen", "")}
    return crear(db, fila, por)


def eliminar(db: Session, id_: int) -> bool:
    r = db.get(DesignCambioDigital, id_)
    if not r:
        return False
    db.delete(r)
    db.commit()
    return True


def _encabezados(hoja):
    """Fila de encabezados: la que tiene DATE / ORDER / PATIENT... Devuelve ({campo: columna}, número de fila) o None."""
    alias = {"fecha": ("date", "fecha"), "orden": ("order", "orden"), "paciente": ("patient", "paciente"), "centro": ("center", "centro"),
             "doctor": ("doctor",), "producto": ("product", "producto"), "tipo_cambio": ("changes", "cambios", "cambio"),
             "detalle": ("details", "detalle", "detalles"), "origen": ("doc/lab", "origen", "lab/doc")}
    for fila in hoja.iter_rows(min_row=1, max_row=min(hoja.max_row, 15)):
        mapa = {}
        for celda in fila:
            t = _sin_acentos(str(celda.value or "")).strip()
            for campo, nombres in alias.items():
                if t in nombres and campo not in mapa:
                    mapa[campo] = celda.column
        if {"fecha", "orden", "paciente", "centro"} <= set(mapa):
            if "origen" not in mapa and "detalle" in mapa:
                mapa["origen"] = mapa["detalle"] + 1   # el reporte trae DOC/LAB en la columna sin encabezado, a la derecha de Detalle
            return mapa, fila[0].row
    return None


def importar_xlsx(db: Session, datos: bytes, por: str) -> dict:
    """Importa las filas del Excel. Sin duplicados: una fila que ya existe (mismo número de orden, fecha, centro y cambio) no se repite."""
    import openpyxl
    try:
        libro = openpyxl.load_workbook(io.BytesIO(datos), data_only=True, read_only=False)
    except Exception:
        raise ValueError("No se pudo leer el archivo. Sube un Excel (.xlsx).")
    hoja = libro.worksheets[0]
    encontrado = _encabezados(hoja)
    if not encontrado:
        raise ValueError("No encontré los encabezados del reporte (DATE, ORDER, PATIENT, CENTER...).")
    mapa, fila_encabezado = encontrado
    existentes = {(r.orden, r.fecha, r.centro, r.tipo_cambio) for r in db.query(DesignCambioDigital).all()}
    agregados, repetidos, errores = 0, 0, []
    for fila in hoja.iter_rows(min_row=fila_encabezado + 1, max_row=hoja.max_row):
        valores = {c: fila[col - 1].value if col - 1 < len(fila) else None for c, col in mapa.items()}
        if not any(v not in (None, "") for v in valores.values()):
            continue
        try:
            d = {c: normalizar(c, valores.get(c)) for c in CAMPOS if c in mapa}
            d["origen"] = normalizar("origen", valores.get("origen")) if "origen" in mapa else ""
            d["orden"] = _texto(valores.get("orden"), "orden")
            for c in ("paciente", "centro", "doctor", "producto", "tipo_cambio", "detalle"):
                d[c] = _texto(valores.get(c), c) if c in mapa else ""
            if not (d["paciente"] or d["orden"]):
                continue
            clave = (d["orden"], d["fecha"], d["centro"], d["tipo_cambio"])
            if clave in existentes:
                repetidos += 1
                continue
            existentes.add(clave)
            db.add(DesignCambioDigital(creado_por=(por or "")[:150], creado_en=datetime.utcnow(), actualizado_en=datetime.utcnow(), **{c: d.get(c, "") for c in CAMPOS}))
            agregados += 1
        except ValueError as ex:
            errores.append(f"Fila {fila[0].row}: {ex}")
    db.commit()
    return {"agregados": agregados, "repetidos": repetidos, "errores": errores[:20]}


def exportar_csv(db: Session, **filtros) -> str:
    datos = listar(db, limite=200, offset=0, **filtros)
    filas = list(datos["items"])
    total = datos["total"]
    # Trae todas las filas que cumplen el filtro (no solo la primera página)
    if total > len(filas):
        filas = []
        for pos in range(0, total, 200):
            filas += listar(db, limite=200, offset=pos, **filtros)["items"]
    salida = io.StringIO()
    w = csv.writer(salida, delimiter=";")
    w.writerow(["Fecha", "Orden", "Paciente", "Centro", "Doctor", "Producto", "Cambio", "Detalle", "Origen"])
    for f in filas:
        w.writerow([f["fecha"], f["orden"], f["paciente"], f["centro"], f["doctor"], f["producto"], f["tipo_cambio"], f["detalle"], f["origen"]])
    return "﻿" + salida.getvalue()
