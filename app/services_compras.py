"""Lógica de negocio del módulo Compras (inventario de Items) -- réplica de app/services_sst.py."""
from datetime import date, datetime
from io import BytesIO
from openpyxl import load_workbook, Workbook
from sqlalchemy import func, or_
from sqlalchemy.orm import Session, joinedload
from .models import Empleado
from .models_compras import (CompraItem, CompraIngreso, CompraSolicitud, CompraSolicitudLinea,
                             CATEGORIAS_COMPRAS, UNIDADES_CONTEO)


# ---------- Catálogo ----------

def catalogo(db: Session, solo_activos: bool = True, buscar: str = "", categoria: str = "") -> list[CompraItem]:
    q = db.query(CompraItem)
    if solo_activos:
        q = q.filter(CompraItem.activo == 1)
    buscar = buscar.strip()
    if buscar:
        patron = f"%{buscar}%"
        q = q.filter(or_(CompraItem.nombre.ilike(patron), CompraItem.sku.ilike(patron)))
    if categoria:
        q = q.filter(CompraItem.categoria == categoria)
    return q.order_by(CompraItem.orden, CompraItem.nombre).all()


def _categoria_valida(categoria: str) -> str:
    return categoria if categoria in CATEGORIAS_COMPRAS else "Otros"


def crear_item(db: Session, nombre: str, sku: str, categoria: str, unidad_conteo: str,
              presentacion: int) -> CompraItem | None:
    nombre = nombre.strip()
    if not nombre or db.query(CompraItem).filter(CompraItem.nombre == nombre).first():
        return None
    orden = (db.query(CompraItem).count() or 0) + 1
    item = CompraItem(nombre=nombre, sku=sku.strip(), categoria=_categoria_valida(categoria),
                      unidad_conteo=unidad_conteo, presentacion=max(1, presentacion), orden=orden)
    db.add(item)
    db.commit()
    db.refresh(item)
    return item


def editar_item(db: Session, item_id: int, nombre: str, sku: str, categoria: str, unidad_conteo: str,
                presentacion: int) -> None:
    item = db.get(CompraItem, item_id)
    if item:
        item.nombre = nombre.strip() or item.nombre
        item.sku = sku.strip()
        item.categoria = _categoria_valida(categoria)
        item.unidad_conteo = unidad_conteo
        item.presentacion = max(1, presentacion)
        db.commit()


def toggle_item(db: Session, item_id: int) -> None:
    item = db.get(CompraItem, item_id)
    if item:
        item.activo = 0 if item.activo else 1
        db.commit()


# ---------- Carga masiva de items (Excel) ----------

COLUMNAS_IMPORT = ["nombre", "sku", "categoria", "unidad_conteo", "presentacion"]
REQUERIDAS_IMPORT = ["nombre"]


def importar_items_excel(db: Session, contenido: bytes) -> dict:
    """Crea/actualiza items desde un xlsx. Upsert por `nombre`. Mismo patrón que
    excel_import.importar_empleados: errores por fila, no se aborta el archivo completo."""
    wb = load_workbook(BytesIO(contenido), data_only=True)
    ws = wb.active
    encabezados = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    faltantes = [c for c in REQUERIDAS_IMPORT if c not in encabezados]
    if faltantes:
        return {"creados": 0, "actualizados": 0,
                "errores": [f"Faltan columnas obligatorias: {', '.join(faltantes)}. Descarga la plantilla."]}
    idx = {c: encabezados.index(c) for c in COLUMNAS_IMPORT if c in encabezados}

    def val(fila, col):
        i = idx.get(col)
        v = fila[i] if i is not None and i < len(fila) else None
        return str(v).strip() if v not in (None, "") else ""

    creados = actualizados = 0
    errores: list[str] = []
    for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(fila):
            continue
        try:
            nombre = val(fila, "nombre")
            if not nombre:
                raise ValueError("nombre es obligatorio")
            categoria_txt = val(fila, "categoria")
            if categoria_txt and categoria_txt not in CATEGORIAS_COMPRAS:
                raise ValueError(f"categoría inválida: '{categoria_txt}' "
                                f"(usa {', '.join(CATEGORIAS_COMPRAS)})")
            unidad_txt = (val(fila, "unidad_conteo") or "UND").upper()
            if unidad_txt not in UNIDADES_CONTEO:
                raise ValueError(f"unidad_conteo inválida: '{unidad_txt}' (usa {', '.join(UNIDADES_CONTEO)})")
            presentacion_txt = val(fila, "presentacion")
            presentacion = max(1, int(float(presentacion_txt))) if presentacion_txt else 1

            item = db.query(CompraItem).filter(CompraItem.nombre == nombre).first()
            nuevo = item is None
            if nuevo:
                orden = (db.query(CompraItem).count() or 0) + 1
                item = CompraItem(nombre=nombre, orden=orden)
                db.add(item)
            item.sku = val(fila, "sku")
            item.categoria = _categoria_valida(categoria_txt) if categoria_txt else (item.categoria if not nuevo else "Otros")
            item.unidad_conteo = unidad_txt
            item.presentacion = presentacion
            db.flush()
            if nuevo:
                creados += 1
            else:
                actualizados += 1
        except Exception as e:  # noqa: BLE001
            errores.append(f"Fila {n}: {e}")

    db.commit()
    return {"creados": creados, "actualizados": actualizados, "errores": errores}


def generar_plantilla_items() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(COLUMNAS_IMPORT)
    ws.append(["Papel higiénico x12", "ASC-001", "Aseo y Cafetería", "PACK", 12])
    ws.append(["Mouse inalámbrico", "TEC-014", "Tecnología", "UND", 1])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------- Ingresos (pendientes de aprobación) ----------

def _normalizar_unidades(item: CompraItem, cantidad: float, unidad_usada: str) -> float:
    if unidad_usada == "PACK" and item.unidad_conteo == "PACK":
        return cantidad * item.presentacion
    return cantidad


def crear_ingreso(db: Session, user: Empleado, item_id: int, cantidad: float, unidad_usada: str,
                  fecha: date, notas: str = "") -> CompraIngreso | None:
    item = db.get(CompraItem, item_id)
    if not item or cantidad <= 0:
        return None
    ingreso = CompraIngreso(item_id=item_id, cantidad_ingresada=cantidad, unidad_usada=unidad_usada,
                            cantidad_unidades=_normalizar_unidades(item, cantidad, unidad_usada),
                            fecha=fecha, registrado_por_id=user.id, notas=notas.strip())
    db.add(ingreso)
    db.commit()
    db.refresh(ingreso)
    return ingreso


def ingresos_pendientes(db: Session) -> list[CompraIngreso]:
    return (db.query(CompraIngreso).options(joinedload(CompraIngreso.item), joinedload(CompraIngreso.registrado_por))
            .filter(CompraIngreso.estado == "pendiente").order_by(CompraIngreso.creado_en).all())


def aprobar_ingreso(db: Session, ingreso: CompraIngreso, user: Empleado) -> None:
    ingreso.estado = "aprobado"
    ingreso.resuelto_por_id = user.id
    ingreso.resuelto_en = datetime.utcnow()
    db.commit()


def rechazar_ingreso(db: Session, ingreso: CompraIngreso, user: Empleado, motivo: str) -> None:
    ingreso.estado = "rechazado"
    ingreso.resuelto_por_id = user.id
    ingreso.resuelto_en = datetime.utcnow()
    ingreso.motivo_rechazo = motivo.strip()
    db.commit()


# ---------- Solicitudes ----------

def crear_solicitud(db: Session, user: Empleado, lineas: list[dict], notas: str = "") -> CompraSolicitud | None:
    """`lineas`: [{"itemId": int, "cantidad": float, "unidadUsada": "PACK"|"UND"|"GALON"}, ...]"""
    if not lineas:
        return None
    solicitud = CompraSolicitud(solicitante_id=user.id, fecha_solicitud=date.today(), notas=notas.strip())
    db.add(solicitud)
    db.flush()
    lineas_validas = 0
    for l in lineas:
        item = db.get(CompraItem, l["itemId"])
        cantidad = float(l["cantidad"])
        if not item or cantidad <= 0:
            continue
        unidad_usada = l["unidadUsada"]
        db.add(CompraSolicitudLinea(solicitud_id=solicitud.id, item_id=item.id, cantidad_solicitada=cantidad,
                                    unidad_usada=unidad_usada,
                                    cantidad_unidades=_normalizar_unidades(item, cantidad, unidad_usada)))
        lineas_validas += 1
    if lineas_validas == 0:
        db.rollback()
        return None
    db.commit()
    db.refresh(solicitud)
    return solicitud


def solicitudes_pendientes(db: Session) -> list[CompraSolicitud]:
    return (db.query(CompraSolicitud)
            .options(joinedload(CompraSolicitud.solicitante),
                    joinedload(CompraSolicitud.lineas).joinedload(CompraSolicitudLinea.item))
            .filter(CompraSolicitud.estado == "pendiente").order_by(CompraSolicitud.fecha_solicitud).all())


def solicitudes_de(db: Session, empleado_id: int) -> list[CompraSolicitud]:
    return (db.query(CompraSolicitud)
            .options(joinedload(CompraSolicitud.lineas).joinedload(CompraSolicitudLinea.item))
            .filter(CompraSolicitud.solicitante_id == empleado_id)
            .order_by(CompraSolicitud.fecha_solicitud.desc()).all())


def entregar_solicitud(db: Session, solicitud: CompraSolicitud, user: Empleado) -> str | None:
    """Devuelve None si se entregó, o un mensaje de error si no hay stock suficiente para alguna línea."""
    stock = stock_a_fecha(db, date.today())
    for linea in solicitud.lineas:
        disponible = stock.get(linea.item_id, 0.0)
        if disponible < linea.cantidad_unidades:
            return (f"No hay suficiente stock de \"{linea.item.nombre}\": disponible {disponible:g}, "
                    f"solicitado {linea.cantidad_unidades:g}.")
    solicitud.estado = "entregada"
    solicitud.entregado_por_id = user.id
    solicitud.entregado_en = datetime.utcnow()
    db.commit()
    return None


def rechazar_solicitud(db: Session, solicitud: CompraSolicitud, user: Empleado, motivo: str) -> None:
    solicitud.estado = "rechazada"
    solicitud.entregado_por_id = user.id
    solicitud.entregado_en = datetime.utcnow()
    solicitud.motivo_rechazo = motivo.strip()
    db.commit()


# ---------- Stock derivado y reportes ----------

def stock_a_fecha(db: Session, fecha: date) -> dict[int, float]:
    """{item_id: unidades en stock a `fecha` inclusive}. Sin tabla de saldo: se deriva sumando
    ingresos aprobados y restando líneas de solicitudes entregadas, igual que SST."""
    saldo: dict[int, float] = {}
    for item_id, total in (db.query(CompraIngreso.item_id, func.coalesce(func.sum(CompraIngreso.cantidad_unidades), 0.0))
                           .filter(CompraIngreso.estado == "aprobado", CompraIngreso.fecha <= fecha)
                           .group_by(CompraIngreso.item_id).all()):
        saldo[item_id] = saldo.get(item_id, 0.0) + float(total)
    salidas = (db.query(CompraSolicitudLinea.item_id, func.coalesce(func.sum(CompraSolicitudLinea.cantidad_unidades), 0.0))
              .join(CompraSolicitud, CompraSolicitudLinea.solicitud_id == CompraSolicitud.id)
              .filter(CompraSolicitud.estado == "entregada", func.date(CompraSolicitud.entregado_en) <= fecha)
              .group_by(CompraSolicitudLinea.item_id).all())
    for item_id, total in salidas:
        saldo[item_id] = saldo.get(item_id, 0.0) - float(total)
    return saldo


def reporte_inventario(db: Session, fecha: date) -> list[dict]:
    saldo = stock_a_fecha(db, fecha)
    return [{"itemId": it.id, "nombre": it.nombre, "sku": it.sku, "categoria": it.categoria,
            "unidadConteo": it.unidad_conteo, "presentacion": it.presentacion,
            "stockUnidades": saldo.get(it.id, 0.0)}
           for it in catalogo(db, solo_activos=False)]


def reporte_movimientos(db: Session, fecha_desde: date, fecha_hasta: date) -> list[dict]:
    movimientos = []
    ingresos = (db.query(CompraIngreso).options(joinedload(CompraIngreso.item), joinedload(CompraIngreso.registrado_por))
               .filter(CompraIngreso.estado == "aprobado", CompraIngreso.fecha >= fecha_desde,
                       CompraIngreso.fecha <= fecha_hasta).all())
    for i in ingresos:
        movimientos.append({"fecha": i.fecha.isoformat(), "tipo": "Ingreso", "item": i.item.nombre,
                            "cantidadUnidades": i.cantidad_unidades,
                            "persona": i.registrado_por.nombre_completo if i.registrado_por else ""})
    entregas = (db.query(CompraSolicitudLinea).join(CompraSolicitud, CompraSolicitudLinea.solicitud_id == CompraSolicitud.id)
               .options(joinedload(CompraSolicitudLinea.item),
                       joinedload(CompraSolicitudLinea.solicitud).joinedload(CompraSolicitud.solicitante))
               .filter(CompraSolicitud.estado == "entregada", func.date(CompraSolicitud.entregado_en) >= fecha_desde,
                       func.date(CompraSolicitud.entregado_en) <= fecha_hasta).all())
    for l in entregas:
        movimientos.append({"fecha": l.solicitud.entregado_en.date().isoformat(), "tipo": "Salida", "item": l.item.nombre,
                            "cantidadUnidades": -l.cantidad_unidades,
                            "persona": l.solicitud.solicitante.nombre_completo if l.solicitud.solicitante else ""})
    movimientos.sort(key=lambda m: m["fecha"])
    return movimientos
