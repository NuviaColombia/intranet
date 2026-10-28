"""Accounts Payable: lógica de negocio -- ver app/models_pagos.py y docs/wiki/caja_menor.md.

Quién aprueba (APROBADORES_AP) viene literalmente del Excel original: la columna "APROVAL" solo
tenía "Leslie" u "Oscar" como valores. Pedido de Oscar (28-oct-2026): cuando se crea una cuenta por
pagar nueva (PENDIENTE), se avisa a los dos por Zoho Cliq y por notificación push."""
import io
from datetime import date, datetime
from sqlalchemy import func
from sqlalchemy.orm import Session
from .models import Empleado
from .models_pagos import CuentaPorPagar, ESTADO_PENDIENTE, ESTADO_APROBADO, ESTADO_PAGADO, ESTADO_OTRO, ESTADOS_AP
from . import zoho_cliq as cliq


class NoEncontrado(LookupError):
    pass


class Conflicto(ValueError):
    pass


# email -> nombre corto, tal como aparecía en la columna APROVAL del Excel
APROBADORES_AP = {
    "oscaralmanza@nuviasmiles.com": "Oscar",
    "leslie.manotas@nuvia.app": "Leslie",
}


def puede_aprobar(user: Empleado) -> bool:
    return user.rol == "superadmin" or (user.email or "").strip().lower() in APROBADORES_AP


def _nombre_aprobador(user: Empleado) -> str:
    return APROBADORES_AP.get((user.email or "").strip().lower()) or user.nombre_completo


# ---------------- Listar / resumen ----------------

def _fila(c: CuentaPorPagar) -> dict:
    return {
        "id": c.id, "empresa": c.empresa, "fecha": c.fecha.isoformat() if c.fecha else None,
        "proveedorNit": c.proveedor_nit, "proveedorNombre": c.proveedor_nombre, "categoria": c.categoria,
        "detalle": c.detalle, "documento": c.documento, "valor": c.valor, "valorUsd": c.valor_usd,
        "fechaVencimiento": c.fecha_vencimiento.isoformat() if c.fecha_vencimiento else None,
        "estado": c.estado, "estadoOriginal": c.estado_original, "aprobadoPor": c.aprobado_por,
        "fechaAprobacion": c.fecha_aprobacion.isoformat() if c.fecha_aprobacion else None,
        "fechaPago": c.fecha_pago.isoformat() if c.fecha_pago else None, "comentarios": c.comentarios,
        "importadoDeExcel": c.importado_de_excel,
    }


def listar(db: Session, estado: str = "", empresa: str = "", busqueda: str = "") -> list[dict]:
    q = db.query(CuentaPorPagar)
    if estado:
        q = q.filter(CuentaPorPagar.estado == estado)
    if empresa:
        q = q.filter(CuentaPorPagar.empresa == empresa)
    if busqueda:
        like = f"%{busqueda.strip()}%"
        q = q.filter((CuentaPorPagar.proveedor_nombre.ilike(like)) | (CuentaPorPagar.proveedor_nit.ilike(like))
                     | (CuentaPorPagar.detalle.ilike(like)))
    filas = q.order_by(CuentaPorPagar.fecha.desc().nullslast(), CuentaPorPagar.id.desc()).limit(2000).all()
    return [_fila(c) for c in filas]


def resumen_conteos(db: Session, empresa: str = "") -> dict:
    q = db.query(CuentaPorPagar.estado, func.count(CuentaPorPagar.id))
    if empresa:
        q = q.filter(CuentaPorPagar.empresa == empresa)
    conteos = dict(q.group_by(CuentaPorPagar.estado).all())
    return {e: conteos.get(e, 0) for e in ESTADOS_AP} | {"TODOS": sum(conteos.values())}


# ---------------- Crear / aprobar / rechazar / pagar ----------------

def crear(db: Session, actor: Empleado, empresa: str, proveedor_nombre: str, valor: float, proveedor_nit: str = "",
          categoria: str = "", detalle: str = "", documento: str = "", fecha_vencimiento: date | None = None,
          comentarios: str = "") -> CuentaPorPagar:
    if empresa not in ("Smiles", "Design"):
        raise Conflicto("La empresa debe ser Smiles o Design.")
    if not proveedor_nombre.strip():
        raise Conflicto("Falta el nombre del proveedor.")
    if not valor or valor <= 0:
        raise Conflicto("El valor debe ser mayor a cero.")
    c = CuentaPorPagar(empresa=empresa, fecha=date.today(), proveedor_nit=proveedor_nit.strip(),
                       proveedor_nombre=proveedor_nombre.strip().upper(), categoria=categoria.strip(),
                       detalle=detalle.strip(), documento=documento.strip(), valor=float(valor),
                       fecha_vencimiento=fecha_vencimiento, estado=ESTADO_PENDIENTE,
                       comentarios=comentarios.strip(), creado_por_id=actor.id)
    db.add(c)
    db.commit()
    db.refresh(c)
    _notificar_pendiente(db, c, actor)
    return c


def _obtener(db: Session, id: int) -> CuentaPorPagar:
    c = db.get(CuentaPorPagar, id)
    if not c:
        raise NoEncontrado("Cuenta por pagar no encontrada.")
    return c


def aprobar(db: Session, actor: Empleado, id: int) -> CuentaPorPagar:
    c = _obtener(db, id)
    if c.estado != ESTADO_PENDIENTE:
        raise Conflicto("Solo se puede aprobar algo que está Pendiente.")
    c.estado, c.aprobado_por, c.fecha_aprobacion = ESTADO_APROBADO, _nombre_aprobador(actor), date.today()
    db.commit()
    return c


def rechazar(db: Session, actor: Empleado, id: int, motivo: str = "") -> CuentaPorPagar:
    c = _obtener(db, id)
    if c.estado != ESTADO_PENDIENTE:
        raise Conflicto("Solo se puede rechazar algo que está Pendiente.")
    c.estado, c.estado_original = ESTADO_OTRO, "Rejected"
    if motivo.strip():
        c.comentarios = (c.comentarios + "\n" if c.comentarios else "") + f"Rechazado por {_nombre_aprobador(actor)}: {motivo.strip()}"
    db.commit()
    return c


def marcar_pagado(db: Session, actor: Empleado, id: int, fecha_pago: date | None = None) -> CuentaPorPagar:
    c = _obtener(db, id)
    if c.estado != ESTADO_APROBADO:
        raise Conflicto("Solo se puede marcar como pagado algo que ya está Aprobado.")
    c.estado, c.fecha_pago = ESTADO_PAGADO, fecha_pago or date.today()
    db.commit()
    return c


def _notificar_pendiente(db: Session, c: CuentaPorPagar, actor: Empleado) -> None:
    emails = list(APROBADORES_AP.keys())
    empresa_nombre = "Nuvia Smiles" if c.empresa == "Smiles" else "Nuvia Design"
    valor_fmt = f"$ {c.valor:,.0f}".replace(",", ".")
    texto = (f"💳 Nueva cuenta por pagar pendiente de aprobar ({empresa_nombre})\n"
             f"{c.proveedor_nombre} · {valor_fmt}" + (f" · {c.detalle}" if c.detalle else "") +
             f"\nRegistrado por {actor.nombre_completo}.")
    try:
        cliq.enviar_cliq_varios(emails, texto, [cliq.boton_enlace("Ver en Accounts Payable", "/pagos")])
    except Exception as e:  # noqa: BLE001 -- un aviso fallido no debe tumbar la creación del registro
        print(f"[Pagos AP] no se pudo enviar el aviso de Cliq: {e}")
    try:
        from . import services_design_push as push
        ids = [e.id for e in db.query(Empleado.id).filter(func.lower(Empleado.email).in_(emails)).all()]
        if ids:
            push.enviar_en_segundo_plano(ids, "💳 Pago pendiente de aprobar",
                                         f"{c.proveedor_nombre} · {valor_fmt} ({empresa_nombre})", "/pagos", "pago-pendiente")
    except Exception as e:  # noqa: BLE001
        print(f"[Pagos AP] no se pudo enviar el push: {e}")


# ---------------- Importar histórico desde el Excel (AP SMILES / AP DESIGN) ----------------

def _num(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("$", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _usd(v) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    return None   # "#VALUE!" (error de fórmula) o vacío


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _clasificar(state_raw) -> tuple[str, str]:
    """(estado, estado_original). Solo "Paid" y "AP" entran al flujo normal -- todo lo demás (Rejected,
    Postponed, Credit, Quote, "-", vacío...) se guarda como OTRO, tal cual, para no inventar un estado
    que el histórico no tenía."""
    s = (str(state_raw).strip().lower()) if state_raw is not None else ""
    if s == "paid":
        return ESTADO_PAGADO, ""
    if s == "ap":
        return ESTADO_APROBADO, ""
    return ESTADO_OTRO, (str(state_raw).strip() if state_raw is not None else "")


# Misma posición de columnas en AP SMILES y AP DESIGN (headers distintos, mismo orden) -- fila 10 es
# el encabezado real, los datos empiezan en la fila 11 (las 9 filas de arriba son un mini resumen
# "State/Quote/AP/Paid" del propio Excel, no son datos).
_FILA_INICIO = 11
_COL = {"fecha": 0, "tipo_mov": 1, "nit": 2, "proveedor": 3, "categoria": 4, "detalle": 5, "documento": 6,
        "valor": 7, "usd": 8, "vencimiento": 9, "aprobador": 10, "estado": 11, "pago": 12, "comentarios": 13}
_HOJAS = (("AP SMILES", "Smiles"), ("AP DESIGN", "Design"))


def importar_excel(db: Session, actor: Empleado, contenido: bytes) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    resumen = {"leidas": 0, "importadas": 0, "duplicadas": 0, "omitidas": 0, "hojas": []}
    existentes = {(c.empresa, c.fecha, c.proveedor_nit, c.detalle, round(c.valor, 2))
                  for c in db.query(CuentaPorPagar.empresa, CuentaPorPagar.fecha, CuentaPorPagar.proveedor_nit,
                                    CuentaPorPagar.detalle, CuentaPorPagar.valor)}
    for nombre_hoja, empresa in _HOJAS:
        if nombre_hoja not in wb.sheetnames:
            continue
        ws = wb[nombre_hoja]
        leidas = importadas = duplicadas = omitidas = 0
        nuevas = []
        for row in ws.iter_rows(min_row=_FILA_INICIO, values_only=True):
            if not any(v is not None for v in row):
                continue
            leidas += 1
            proveedor = _texto(row[_COL["proveedor"]])
            detalle = _texto(row[_COL["detalle"]])
            valor = _num(row[_COL["valor"]])
            if not proveedor and not detalle:
                omitidas += 1
                continue
            fecha = _fecha(row[_COL["fecha"]])
            nit = _texto(row[_COL["nit"]])
            clave = (empresa, fecha, nit, detalle, round(valor, 2))
            if clave in existentes:
                duplicadas += 1
                continue
            existentes.add(clave)
            estado, estado_original = _clasificar(row[_COL["estado"]])
            tipo_mov = _texto(row[_COL["tipo_mov"]])
            comentarios = _texto(row[_COL["comentarios"]])
            if tipo_mov and tipo_mov != "For Payment":
                comentarios = f"[{tipo_mov}] " + comentarios
            nuevas.append(CuentaPorPagar(
                empresa=empresa, fecha=fecha, proveedor_nit=nit, proveedor_nombre=proveedor,
                categoria=_texto(row[_COL["categoria"]]), detalle=detalle, documento=_texto(row[_COL["documento"]]),
                valor=valor, valor_usd=_usd(row[_COL["usd"]]), fecha_vencimiento=_fecha(row[_COL["vencimiento"]]),
                estado=estado, estado_original=estado_original, aprobado_por=_texto(row[_COL["aprobador"]]),
                fecha_pago=_fecha(row[_COL["pago"]]) if estado == ESTADO_PAGADO else None,
                comentarios=comentarios, creado_por_id=actor.id, importado_de_excel=True,
            ))
        db.add_all(nuevas)
        db.commit()
        importadas = len(nuevas)
        resumen["leidas"] += leidas
        resumen["importadas"] += importadas
        resumen["duplicadas"] += duplicadas
        resumen["omitidas"] += omitidas
        resumen["hojas"].append({"hoja": nombre_hoja, "empresa": empresa, "leidas": leidas,
                                 "importadas": importadas, "duplicadas": duplicadas, "omitidas": omitidas})
    return resumen
