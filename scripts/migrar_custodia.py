"""Migra el historial de 'Cambio de Custodia' desde el Excel exportado de Google Sheets.

Agrupa las filas de REGISTRO por CONSECUTIVO en un CustodiaTraslado (cabecera) + N
CustodiaOrdenLinea (una por sub-orden), preservando el CONSECUTIVO original como id.
Enlaza RESUMEN_GENERAL/DETALLE_DISCOS/DETALLE_OP por ese mismo consecutivo.

Uso: apunta DATABASE_URL a la base de destino (ver .env) y ejecuta:
    python -m scripts.migrar_custodia "ruta/al/CAMBIO DE CUSTODIA 2.xlsx"
"""
import sys
from datetime import datetime, date, time, timedelta
from openpyxl import load_workbook
from sqlalchemy import text
from app.database import SessionLocal, engine
from app.models import Empleado
from app.models_custodia import CustodiaTraslado, CustodiaOrdenLinea, CustodiaResumen, CustodiaDiscos, CustodiaOP


def _excel_serial_a_datetime(valor):
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, (int, float)):
        return datetime(1899, 12, 30) + timedelta(days=valor)
    return None


def _texto(valor) -> str:
    """Convierte una celda a texto, evitando el '.0' que Excel/openpyxl agrega a
    números enteros almacenados como float (ids, números de orden, códigos de OP)."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def migrar(ruta_xlsx: str):
    wb = load_workbook(ruta_xlsx, data_only=True)
    ws_reg = wb["REGISTRO"]

    db = SessionLocal()
    emails_conocidos = {e.email.lower(): e.id for e in db.query(Empleado).all()}

    existentes = {i for (i,) in db.query(CustodiaTraslado.id).all()}
    if existentes:
        print(f"Ya existen {len(existentes)} traslados en la base de destino; se omitirán esos consecutivos.")

    grupos: dict[int, dict] = {}
    filas_omitidas = 0
    for n, fila in enumerate(ws_reg.iter_rows(min_row=2, values_only=True), start=2):
        if not any(fila):
            continue
        crudo = fila[13]
        if crudo is None:
            filas_omitidas += 1
            continue
        consecutivo = int(crudo)
        if consecutivo in existentes:
            continue

        numero_orden = _texto(fila[3])
        cantidad = float(fila[9] or 0)

        if consecutivo not in grupos:
            motivo = _texto(fila[10]).upper()
            anulado = motivo == "ANULADO"
            anulado_en = _excel_serial_a_datetime(fila[11]) if anulado else None
            usuario_anula = _texto(fila[12])
            anulado_por_id = emails_conocidos.get(usuario_anula.lower())

            estado_firmas = _texto(fila[14]).upper()
            sello_entrada = _texto(fila[16])
            # Los movimientos históricos ya ocurrieron: se marcan como confirmados por defecto
            # (no tiene sentido dejar cientos de registros antiguos "pendientes de confirmar").
            confirmado_entrada = True

            fecha_valor = fila[4]
            fecha = fecha_valor.date() if isinstance(fecha_valor, datetime) else (fecha_valor or date(2000, 1, 1))
            hora_valor = fila[5]
            if isinstance(hora_valor, datetime):
                hora = hora_valor.time()
            elif isinstance(hora_valor, time):
                hora = hora_valor
            else:
                hora = time(0, 0)

            grupos[consecutivo] = {
                "colaborador": _texto(fila[0]) or "(sin nombre)",
                "id_colaborador": _texto(fila[1]),
                "area_creacion": _texto(fila[2]).upper(),
                "fecha": fecha, "hora": hora, "usuario": _texto(fila[6]),
                "area_salida": _texto(fila[7]).upper(), "area_entrada": _texto(fila[8]).upper(),
                "motivo": motivo or "PRODUCCIÓN NORMAL",
                "anulado": anulado, "anulado_en": anulado_en, "anulado_por_id": anulado_por_id,
                "confirmado_entrada": confirmado_entrada,
                "lineas": [],
            }
        if numero_orden:
            grupos[consecutivo]["lineas"].append((numero_orden.upper(), cantidad))

    print(f"REGISTRO: {ws_reg.max_row - 1} filas -> {len(grupos)} traslados nuevos "
         f"({filas_omitidas} filas sin CONSECUTIVO omitidas).")

    for consecutivo, datos in grupos.items():
        db.add(CustodiaTraslado(
            id=consecutivo, colaborador=datos["colaborador"], id_colaborador=datos["id_colaborador"],
            area_creacion=datos["area_creacion"], fecha=datos["fecha"], hora=datos["hora"],
            usuario=datos["usuario"], area_salida=datos["area_salida"], area_entrada=datos["area_entrada"],
            motivo=datos["motivo"], creado_por_id=None, confirmado_entrada=datos["confirmado_entrada"],
            anulado=datos["anulado"], anulado_en=datos["anulado_en"], anulado_por_id=datos["anulado_por_id"],
        ))
        for numero_orden, cantidad in datos["lineas"]:
            db.add(CustodiaOrdenLinea(traslado_id=consecutivo, numero_orden=numero_orden, cantidad_discos=cantidad))
    db.commit()
    print(f"{len(grupos)} traslados y sus líneas guardados.")

    def _migrar_detalle(nombre_hoja, modelo, mapeador):
        if nombre_hoja not in wb.sheetnames:
            return 0
        total = 0
        for fila in wb[nombre_hoja].iter_rows(min_row=2, values_only=True):
            if not any(fila) or fila[0] is None:
                continue
            consecutivo = int(fila[0])
            if consecutivo not in grupos:
                continue
            db.add(modelo(traslado_id=consecutivo, **mapeador(fila)))
            total += 1
        return total

    n_res = _migrar_detalle("RESUMEN_GENERAL", CustodiaResumen, lambda f: {
        "orden": _texto(f[1]), "descripcion": _texto(f[2]), "paciente": _texto(f[3]), "total": f[4] or 0,
        "verificado_salida": _texto(f[5]).upper() == "SI", "verificado_entrada": _texto(f[6]).upper() == "SI",
    })
    n_disc = _migrar_detalle("DETALLE_DISCOS", CustodiaDiscos, lambda f: {
        "detalle_protesis": _texto(f[1]), "cant_paciente": f[2] or 0, "cant_discos": f[3] or 0,
    })
    n_op = _migrar_detalle("DETALLE_OP", CustodiaOP, lambda f: {
        "orden": _texto(f[3]), "op": _texto(f[1]), "descripcion": _texto(f[2]), "tipo": _texto(f[4]),
        "usuario": _texto(f[5]), "observaciones": _texto(f[6]),
    })
    db.commit()
    print(f"Detalle migrado: resumen={n_res} discos={n_disc} op={n_op}")

    if engine.dialect.name == "postgresql":
        with engine.begin() as conn:
            conn.execute(text(
                "SELECT setval(pg_get_serial_sequence('custodia_traslados', 'id'), "
                "(SELECT COALESCE(MAX(id), 1) FROM custodia_traslados))"))
        print("Secuencia de Postgres reiniciada para que los próximos registros continúen tras el máximo migrado.")

    db.close()


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Uso: python -m scripts.migrar_custodia <ruta al xlsx>")
        sys.exit(1)
    migrar(sys.argv[1])
