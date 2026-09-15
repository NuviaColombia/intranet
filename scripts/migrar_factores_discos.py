"""Carga/actualiza el catálogo de tipos de prótesis y su factor de conversión a discos,
leyendo el rango E3:G50 de la hoja DATOS exportada (misma ubicación que usaba
obtenerFactoresDiscos() en el Apps Script original).

Uso: apunta DATABASE_URL a la base de destino (ver .env) y ejecuta:
    python -m scripts.migrar_factores_discos "ruta/al/archivo.xlsx"
"""
import sys
from openpyxl import load_workbook
from app.database import SessionLocal
from app.models_custodia import CustodiaFactorDisco


def migrar(ruta_xlsx: str):
    wb = load_workbook(ruta_xlsx, data_only=True)
    if "DATOS" not in wb.sheetnames:
        print("No se encontró una hoja 'DATOS' en el archivo.")
        sys.exit(1)
    ws = wb["DATOS"]

    db = SessionLocal()
    existentes = {f.detalle: f for f in db.query(CustodiaFactorDisco).all()}

    creados, actualizados, orden = 0, 0, 0
    for fila in ws["E3:G50"]:
        detalle_celda, _cant_paciente, factor_celda = fila
        detalle = str(detalle_celda.value).strip() if detalle_celda.value is not None else ""
        if not detalle or detalle.upper() == "TOTAL":
            continue
        factor_raw = str(factor_celda.value).strip().replace(",", ".") if factor_celda.value is not None else ""
        try:
            factor = float(factor_raw)
        except ValueError:
            continue

        orden += 1
        if detalle in existentes:
            existentes[detalle].factor = factor
            existentes[detalle].orden = orden
            actualizados += 1
        else:
            db.add(CustodiaFactorDisco(detalle=detalle, factor=factor, orden=orden))
            creados += 1

    db.commit()
    print(f"Catálogo de discos: {creados} creados, {actualizados} actualizados.")
    db.close()


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Uso: python -m scripts.migrar_factores_discos <ruta al xlsx>")
        sys.exit(1)
    migrar(sys.argv[1])
